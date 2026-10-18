"""Safe CSV/XLSX readers and explicit conversion to canonical eye-tracking tables."""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
import hashlib
from pathlib import Path
from pathlib import PurePosixPath, PureWindowsPath
import re
from typing import Any, Literal, Mapping, Sequence
import zipfile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .contracts import (
    AOI_REQUIRED_COLUMNS,
    CONTRACT_VERSION,
    SEGMENT_REQUIRED_COLUMNS,
    SIGNAL_DERIVED_COLUMNS,
    SIGNAL_OPTIONAL_COLUMNS,
    SIGNAL_REQUIRED_COLUMNS,
    CoordinateMetadata,
    ContractError,
    NormalizedDataset,
)
from .aoi_geometry import (
    AOIGeometryError,
    AOIGeometryIssue,
    AOITableGeometryError,
    AOI_DOMAIN_SCHEMA_VERSION,
    AOI_GEOMETRY_VERSION,
    GeometrySource,
    prepare_static_aois,
)
from .scopes import SegmentScopeConfig, ScopeError, _observation_identity_frame_owned, normalize_segment_records

MAX_IMPORT_BYTES = 2 * 1024**3
MAX_XLSX_UNCOMPRESSED_BYTES = 8 * 1024**3
MAX_XLSX_MEMBERS = 10_000
MAX_XLSX_COMPRESSION_RATIO = 200
SUPPORTED_SUFFIXES = {".csv", ".xlsx"}


class DataImportError(ValueError):
    """An input could not be safely or unambiguously imported."""


class AOIGeometryImportError(DataImportError):
    """Structured AOI source-table error with actual import provenance."""

    def __init__(self, issues: Sequence[AOIGeometryIssue | Mapping[str, Any]]):
        self.issues = [issue.as_dict() if isinstance(issue, AOIGeometryIssue) else dict(issue)
                       for issue in issues]
        first = self.issues[0] if self.issues else {"code": "aoi_import_invalid"}
        super().__init__(f"AOI table rejected ({first.get('code')}): {first.get('message', 'invalid AOI source')}")


@dataclass(frozen=True)
class CsvOptions:
    encoding: str = "utf-8-sig"
    delimiter: str | None = ","
    decimal: str = "."

    def __post_init__(self) -> None:
        if not self.encoding.strip():
            raise DataImportError("CSV encoding must be selected explicitly")
        if self.delimiter is not None and len(self.delimiter) != 1:
            raise DataImportError("CSV delimiter must be one character or None for detection")
        if len(self.decimal) != 1:
            raise DataImportError("CSV decimal separator must be one character")


@dataclass(frozen=True)
class ImportedTable:
    role: Literal["signals", "segments", "aois", "performance"]
    frame: pd.DataFrame
    source_path: Path
    source_name: str
    sha256: str
    byte_size: int
    sheet_name: str | None = None
    source_rows: tuple[int, ...] = ()

    def manifest_entry(self) -> dict[str, Any]:
        return {
            "role": self.role,
            "source_name": self.source_name,
            "sha256": self.sha256,
            "byte_size": self.byte_size,
            "format": self.source_path.suffix.lower().lstrip("."),
            "sheet_name": self.sheet_name,
            "row_count": int(len(self.frame)),
        }


@dataclass(frozen=True)
class SignalImportOptions:
    column_map: Mapping[str, str] = field(default_factory=dict)
    subject_id: str | None = None
    session_id: str | None = None
    time: CoordinateMetadata = field(default_factory=CoordinateMetadata)
    validity_values: Mapping[str, Mapping[str, bool]] = field(default_factory=dict)
    reading_direction_values: Mapping[str, str] = field(default_factory=dict)

    def __post_init__(self) -> None:
        allowed = set(SIGNAL_REQUIRED_COLUMNS) | set(SIGNAL_OPTIONAL_COLUMNS) | {"timestamp"}
        unknown = set(self.column_map) - allowed
        if unknown:
            raise DataImportError("Unsupported canonical signal fields: " + ", ".join(sorted(unknown)))
        boolean_fields = {
            "gaze_valid", "pupil_left_valid", "pupil_right_valid", "gaze_left_valid",
            "gaze_right_valid", "tracking_valid", "target_valid", "reference_valid",
            "eyelid_closed_left", "eyelid_closed_right",
        }
        unknown_validity = set(self.validity_values) - boolean_fields
        if unknown_validity:
            raise DataImportError("Unsupported validity fields: " + ", ".join(sorted(unknown_validity)))
        for name, values in self.validity_values.items():
            if not isinstance(values, Mapping) or any(not isinstance(value, bool) for value in values.values()):
                raise DataImportError(f"Custom validity codes for {name!r} must map labels to booleans")
        if any(not isinstance(canonical, str) or not isinstance(source, str) or not source.strip()
               for canonical, source in self.column_map.items()):
            raise DataImportError("Column mappings must use nonempty string source names")
        if any(not isinstance(source, str) or not isinstance(direction, str)
               or direction.strip().casefold() not in READING_DIRECTIONS
               for source, direction in self.reading_direction_values.items()):
            raise DataImportError("Reading-direction codes must map to a supported canonical direction")
        if self.subject_id is not None and not str(self.subject_id).strip():
            raise DataImportError("Explicit subject_id cannot be blank")
        if self.session_id is not None and not str(self.session_id).strip():
            raise DataImportError("Explicit session_id cannot be blank")


READING_DIRECTIONS = frozenset({"left_to_right", "right_to_left", "top_to_bottom", "bottom_to_top"})


COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "subject_id": ("subject", "participant", "participant_id", "被试", "被试编号", "参与者", "参与者编号"),
    "session_id": ("session", "recording_id", "采集段", "会话", "会话编号"),
    "interaction_id": ("interaction", "team_id", "互动单元", "交互单元"),
    "timestamp": ("time", "timestamp", "time_stamp", "时间", "时间戳", "采样时间"),
    "time_s": ("time_s", "timestamp_s", "time_seconds", "时间秒"),
    "gaze_x": ("x", "gaze_x", "gaze_x_px", "注视点x", "注视x"),
    "gaze_y": ("y", "gaze_y", "gaze_y_px", "注视点y", "注视y"),
    "gaze_z": ("z", "gaze_z", "注视点z", "注视z"),
    "gaze_left_x": ("gaze_left_x", "left_gaze_x", "left_gaze_point_x", "左眼注视x"),
    "gaze_left_y": ("gaze_left_y", "left_gaze_y", "left_gaze_point_y", "左眼注视y"),
    "gaze_left_z": ("gaze_left_z", "left_gaze_z", "left_gaze_point_z", "左眼注视z"),
    "gaze_right_x": ("gaze_right_x", "right_gaze_x", "right_gaze_point_x", "右眼注视x"),
    "gaze_right_y": ("gaze_right_y", "right_gaze_y", "right_gaze_point_y", "右眼注视y"),
    "gaze_right_z": ("gaze_right_z", "right_gaze_z", "right_gaze_point_z", "右眼注视z"),
    "pupil_left": ("left_pupil", "pupil_left", "left_pupil_diameter", "左瞳孔", "左眼瞳孔"),
    "pupil_right": ("right_pupil", "pupil_right", "right_pupil_diameter", "右瞳孔", "右眼瞳孔"),
    "gaze_valid": ("gaze_valid", "gaze_validity", "validity", "注视有效性"),
    "pupil_left_valid": ("pupil_left_valid", "pupil_validity_left", "左眼有效性"),
    "pupil_right_valid": ("pupil_right_valid", "pupil_validity_right", "右眼有效性"),
    "gaze_left_valid": ("gaze_left_valid", "left_gaze_valid", "左眼注视有效性"),
    "gaze_right_valid": ("gaze_right_valid", "right_gaze_valid", "右眼注视有效性"),
    "tracking_valid": ("tracking_valid", "track_valid", "追踪有效"),
    "eyelid_closure_left": ("eyelid_closure_left", "left_eyelid_closure", "左眼闭合"),
    "eyelid_closure_right": ("eyelid_closure_right", "right_eyelid_closure", "右眼闭合"),
    "blink_left_event": ("blink_left_event", "left_blink", "左眼眨眼事件"),
    "blink_right_event": ("blink_right_event", "right_blink", "右眼眨眼事件"),
    "microsaccade_event": ("microsaccade_event", "micro_saccade", "微扫视事件"),
    "tracking_state": ("tracking_state", "track_state", "追踪状态"),
    "calibration_error": ("calibration_error", "accuracy_error", "校准误差"),
    "target_vx": ("target_vx", "target_velocity_x", "目标速度x"),
    "target_vy": ("target_vy", "target_velocity_y", "目标速度y"),
    "target_vz": ("target_vz", "target_velocity_z", "目标速度z"),
    "target_azimuth": ("target_azimuth", "target_angle_horizontal", "目标方位角"),
    "target_elevation": ("target_elevation", "target_angle_vertical", "目标俯仰角"),
    "target_distance": ("target_distance", "target_depth", "目标距离"),
    "target_valid": ("target_valid", "目标有效性"),
    "cue_id": ("cue_id", "cue", "提示编号"),
    "cue_start": ("cue_start", "cue_onset", "提示开始"),
    "cue_end": ("cue_end", "cue_offset", "提示结束"),
    "stimulus_id": ("stimulus", "stimulus_id", "刺激", "刺激编号"),
    "condition": ("condition", "条件"),
    "trial": ("trial", "trial_id", "试次", "试次编号"),
}
COLUMN_ALIASES.update({
    "gaze_left_direction_x": ("left_gaze_direction_x", "left_direction_x", "左眼视线方向x"),
    "gaze_left_direction_y": ("left_gaze_direction_y", "left_direction_y", "左眼视线方向y"),
    "gaze_left_direction_z": ("left_gaze_direction_z", "left_direction_z", "左眼视线方向z"),
    "gaze_right_direction_x": ("right_gaze_direction_x", "right_direction_x", "右眼视线方向x"),
    "gaze_right_direction_y": ("right_gaze_direction_y", "right_direction_y", "右眼视线方向y"),
    "gaze_right_direction_z": ("right_gaze_direction_z", "right_direction_z", "右眼视线方向z"),
    "gaze_left_origin_x": ("left_eye_origin_x", "left_origin_x", "左眼原点x"),
    "gaze_left_origin_y": ("left_eye_origin_y", "left_origin_y", "左眼原点y"),
    "gaze_left_origin_z": ("left_eye_origin_z", "left_origin_z", "左眼原点z"),
    "gaze_right_origin_x": ("right_eye_origin_x", "right_origin_x", "右眼原点x"),
    "gaze_right_origin_y": ("right_eye_origin_y", "right_origin_y", "右眼原点y"),
    "gaze_right_origin_z": ("right_eye_origin_z", "right_origin_z", "右眼原点z"),
    "eyelid_closed_left": ("left_eyelid_closed", "左眼闭合状态"),
    "eyelid_closed_right": ("right_eyelid_closed", "右眼闭合状态"),
    "target_id": ("target_id", "target", "目标编号"),
    "reference_id": ("reference_id", "reference", "参考编号"),
    "reference_x": ("reference_x", "reference_point_x", "参考点x"),
    "reference_y": ("reference_y", "reference_point_y", "参考点y"),
    "reference_z": ("reference_z", "reference_point_z", "参考点z"),
    "reference_vx": ("reference_vx", "reference_velocity_x", "参考速度x"),
    "reference_vy": ("reference_vy", "reference_velocity_y", "参考速度y"),
    "reference_vz": ("reference_vz", "reference_velocity_z", "参考速度z"),
    "reference_valid": ("reference_valid", "reference_validity", "参考有效性"),
    "reading_direction": ("reading_direction", "reading direction", "阅读方向"),
    "reading_scope_id": ("reading_scope_id", "line_id", "阅读范围编号"),
    "cue_time": ("cue_time", "cue_onset_time", "提示时刻"),
    "cue_type": ("cue_type", "提示类型"),
    "reference_time": ("reference_time", "reference_onset_time", "参考时刻"),
    "reference_event_id": ("reference_event_id", "参考事件编号"),
    "clock_id": ("clock_id", "clock_source_id", "时钟编号"),
    "frame_id": ("frame_id", "video_frame_id", "帧编号"),
    "frame_number": ("frame_number", "frame_index", "帧序号"),
    "timecode": ("timecode", "time_code", "时间码"),
    "device_validation_status": ("device_validation_status", "设备验证状态"),
    "device_validation_method": ("device_validation_method", "设备验证方法"),
    "device_validation_evidence": ("device_validation_evidence", "设备验证依据"),
})


def read_table(
    path: str | Path,
    *,
    role: Literal["signals", "segments", "aois", "performance"],
    sheet_name: str | None = None,
    csv_options: CsvOptions | None = None,
    preview_rows: int | None = None,
) -> ImportedTable:
    """Read a supported local file without executing macros or formulas."""
    source = Path(path).expanduser().resolve(strict=True)
    if not source.is_file():
        raise DataImportError("Input path is not a regular file")
    if source.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise DataImportError("Only CSV and XLSX are supported; XLSM/macros are rejected")
    size = source.stat().st_size
    if size > MAX_IMPORT_BYTES:
        raise DataImportError(f"Input exceeds the bounded {MAX_IMPORT_BYTES}-byte import limit")
    digest = _sha256_file(source)
    if source.suffix.lower() == ".csv":
        frame = _read_csv(source, csv_options or CsvOptions(), preview_rows)
        selected_sheet = None
        source_rows = tuple(range(2, len(frame) + 2))
    else:
        frame, selected_sheet, source_rows = _read_xlsx(source, sheet_name, preview_rows)
    return ImportedTable(
        role=role,
        frame=frame,
        source_path=source,
        source_name=source.name,
        sha256=digest,
        byte_size=size,
        sheet_name=selected_sheet,
        source_rows=source_rows,
    )


def normalize_signals(imported: ImportedTable, options: SignalImportOptions) -> NormalizedDataset:
    """Create the canonical long signal table using only explicit time/ID semantics."""
    if imported.role != "signals":
        raise DataImportError("normalize_signals requires a signals table")
    raw = imported.frame
    if raw.columns.has_duplicates:
        raise DataImportError("Duplicate source headers are ambiguous; rename or map them explicitly")
    mapping = _resolve_mapping(
        raw.columns,
        options.column_map,
        (),
        SIGNAL_OPTIONAL_COLUMNS + ("subject_id", "session_id", "time_s", "timestamp", "original_index"),
    )
    time_sources = list(dict.fromkeys(
        source for source in (mapping.get("time_s"), mapping.get("timestamp")) if source is not None
    ))
    if len(time_sources) > 1:
        raise DataImportError("More than one time column is plausible; map the intended source explicitly")
    time_source = time_sources[0] if time_sources else None
    if time_source is None:
        raise DataImportError("A timestamp column is required")
    metadata = options.time
    if metadata.time_unit == "unknown":
        raise DataImportError("Confirm time_unit as s, ms, frame, or datetime before normalization")

    out = pd.DataFrame(index=raw.index)
    if "original_index" in mapping:
        out["original_index"] = raw[mapping["original_index"]]
    else:
        out["original_index"] = list(imported.source_rows or range(len(raw)))
    for canonical in ("subject_id", "session_id", "interaction_id", "stimulus_id", "condition", "trial"):
        source_column = mapping.get(canonical)
        if source_column is not None:
            out[canonical] = raw[source_column]
    manual_mappings: dict[str, dict[str, Any]] = {}
    for canonical, explicit in (("subject_id", options.subject_id), ("session_id", options.session_id)):
        if canonical not in out or out[canonical].isna().all() or out[canonical].astype("string").str.strip().eq("").all():
            if explicit is None:
                raise DataImportError(f"{canonical} is absent; assign it explicitly for this import batch")
            out[canonical] = str(explicit).strip()
            manual_mappings[canonical] = {
                "operation": "user_confirmed_constant",
                "value": str(explicit).strip(),
                "source_column": mapping.get(canonical),
            }
        else:
            out[canonical] = out[canonical].astype("string").str.strip()
            blanks = out[canonical].isna() | out[canonical].eq("")
            if blanks.any():
                if explicit is None:
                    raise DataImportError(f"Some {canonical} values are blank; no ID will be inferred")
                out.loc[blanks, canonical] = str(explicit).strip()
                manual_mappings[canonical] = {
                    "operation": "user_confirmed_constant_fill",
                    "value": str(explicit).strip(),
                    "source_column": mapping.get(canonical),
                    "filled_rows": int(blanks.sum()),
                }
    out["time_raw"] = raw[time_source]
    out["time_s"], time_parse_error = _convert_time(raw[time_source], metadata)
    out["time_valid"] = out["time_s"].notna()

    raw_validity_names: dict[str, str] = {}
    numeric_parse_errors: dict[str, int] = {}
    numeric_nonfinite_values: dict[str, int] = {}
    numeric_columns = {
        "gaze_x", "gaze_y", "gaze_z", "gaze_left_x", "gaze_left_y", "gaze_left_z",
        "gaze_right_x", "gaze_right_y", "gaze_right_z", "pupil_left", "pupil_right",
        "confidence_left", "confidence_right", "target_x", "target_y", "target_z",
        "target_vx", "target_vy", "target_vz", "target_azimuth", "target_elevation",
        "target_distance", "calibration_error", "frame_number",
        "gaze_left_direction_x", "gaze_left_direction_y", "gaze_left_direction_z",
        "gaze_right_direction_x", "gaze_right_direction_y", "gaze_right_direction_z",
        "gaze_left_origin_x", "gaze_left_origin_y", "gaze_left_origin_z",
        "gaze_right_origin_x", "gaze_right_origin_y", "gaze_right_origin_z",
        "eyelid_closure_left", "eyelid_closure_right",
        "reference_x", "reference_y", "reference_z", "reference_vx", "reference_vy", "reference_vz",
    }
    validity_columns = {
        "gaze_valid", "pupil_left_valid", "pupil_right_valid", "gaze_left_valid",
        "gaze_right_valid", "tracking_valid", "target_valid", "reference_valid",
        "eyelid_closed_left", "eyelid_closed_right",
    }
    for canonical in SIGNAL_OPTIONAL_COLUMNS:
        if canonical in {"interaction_id", "stimulus_id", "condition", "trial"}:
            continue
        source_column = mapping.get(canonical)
        if source_column is None:
            out[canonical] = pd.NA
            continue
        source_values = raw[source_column]
        if canonical in validity_columns:
            raw_name = "source_" + canonical
            out[raw_name] = source_values
            raw_validity_names[canonical] = raw_name
            out[canonical] = _decode_validity(
                source_values, options.validity_values.get(canonical, {})
            )
        elif canonical == "reading_direction":
            out[canonical] = _decode_reading_direction(source_values, options.reading_direction_values)
        elif canonical in numeric_columns:
            numeric = pd.to_numeric(source_values, errors="coerce")
            bad = source_values.notna() & source_values.astype("string").str.strip().ne("") & numeric.isna()
            numeric, nonfinite = _finite_numeric_series(numeric)
            out[canonical] = numeric
            out["parse_error__" + canonical] = bad.astype("boolean")
            if nonfinite.any():
                out["nonfinite__" + canonical] = nonfinite
                numeric_nonfinite_values[canonical] = int(nonfinite.sum())
            numeric_parse_errors[canonical] = int(bad.sum())
        else:
            out[canonical] = source_values

    source_id = f"source:{imported.sha256}"
    out["source_id"] = source_id
    if "timebase_id" not in out:
        out["timebase_id"] = pd.NA
    if metadata.clock_source or metadata.frame_id:
        declared_timebase = metadata.clock_source or metadata.frame_id
        out["timebase_id"] = out["timebase_id"].where(
            out["timebase_id"].notna() & out["timebase_id"].astype("string").str.strip().ne(""),
            str(declared_timebase),
        )
    try:
        out, _ = _observation_identity_frame_owned(out, source_id=source_id)
    except ScopeError as exc:
        raise DataImportError(str(exc)) from exc

    # Cue/reference annotations are retained as source values and also normalized
    # through the exact same confirmed clock conversion used by the signal axis.
    event_time_parse_errors: dict[str, int] = {}
    for canonical in ("cue_time", "cue_start", "cue_end", "reference_time"):
        source_column = mapping.get(canonical)
        if source_column is None:
            out[canonical + "_s"] = pd.NA
            event_time_parse_errors[canonical] = 0
        else:
            seconds, parse_error = _convert_time(raw[source_column], metadata)
            out[canonical + "_s"] = seconds
            out[canonical + "_time_parse_error"] = parse_error.astype("boolean")
            event_time_parse_errors[canonical] = int(parse_error.sum())
    for derived in SIGNAL_DERIVED_COLUMNS:
        if derived not in out:
            out[derived] = pd.NA

    # Ensure the contract's optional columns exist, but missing values stay unknown.
    for canonical in SIGNAL_OPTIONAL_COLUMNS:
        if canonical not in out:
            out[canonical] = pd.NA
    out["time_parse_error"] = time_parse_error.astype("boolean")
    out["source_row"] = out["original_index"]
    for name in validity_columns:
        out[name] = out[name].astype("boolean")

    # Keep every imported field losslessly alongside its normalized representation,
    # including vendor extensions not yet assigned scientific semantics.
    raw_source_columns: dict[str, str] = {}
    used_names = set(out.columns)
    for source_name in raw.columns:
        raw_column = "raw_source__" + str(source_name)
        suffix = 2
        while raw_column in used_names:
            raw_column = f"raw_source__{source_name}__{suffix}"
            suffix += 1
        used_names.add(raw_column)
        out[raw_column] = raw[source_name].copy(deep=True)
        raw_source_columns[str(source_name)] = raw_column

    source_manifest = [imported.manifest_entry()]
    normalized_metadata = {
        "contract_version": CONTRACT_VERSION,
        "source_manifest": source_manifest,
        "column_mapping": dict(mapping),
        "coordinate_metadata": metadata.to_dict(),
        "validity_raw_columns": raw_validity_names,
        "raw_source_columns": raw_source_columns,
        "extension_field_semantics": {
            canonical: {
                "source_column": source_column,
                "unit": "unknown",
                "interpretation": "preserved_without_assuming_vendor_semantics",
            }
            for canonical, source_column in mapping.items()
            if canonical in SIGNAL_OPTIONAL_COLUMNS and canonical not in {
                "condition", "trial", "interaction_id", "stimulus_id", "device_id", "sample_index",
                "vendor_event", "eyelid", "cue_id", "tracking_state", "microsaccade_event",
                "blink_left_event", "blink_right_event",
            }
        },
        "numeric_parse_errors": numeric_parse_errors,
        "numeric_nonfinite_values": numeric_nonfinite_values,
        "time_parse_errors": int(time_parse_error.sum()),
        "event_time_parse_errors": event_time_parse_errors,
        "manual_mappings": manual_mappings,
        "imported_at_utc": datetime.now(timezone.utc).isoformat(),
        "raw_source_retained": True,
    }
    out.index = pd.RangeIndex(len(out))
    dataset = NormalizedDataset(signals=out, metadata=normalized_metadata,
                                source_manifest=source_manifest)
    try:
        dataset.validate()
    except ContractError as exc:
        raise DataImportError(str(exc)) from exc
    return dataset


def normalize_segments(
    imported: ImportedTable,
    *,
    column_map: Mapping[str, str] | None = None,
    time: CoordinateMetadata,
    subject_id: str | None = None,
    session_id: str | None = None,
    scope_config: SegmentScopeConfig | Mapping[str, Any] | None = None,
    member_session_map: Mapping[str, Any] | Sequence[Any] | None = None,
    dataset_scope: Any = None,
    timebase_id: str | None = None,
) -> pd.DataFrame:
    if imported.role != "segments":
        raise DataImportError("normalize_segments requires a segments table")
    raw = imported.frame
    if isinstance(scope_config, SegmentScopeConfig):
        scope_mode = scope_config.mode
    elif isinstance(scope_config, Mapping):
        scope_mode = str(scope_config.get("mode", "participant_session"))
    else:
        scope_mode = "participant_session"
    if scope_mode not in {"participant_session", "shared_interaction"}:
        raise DataImportError("Unsupported segment scope mode")
    requested = dict(column_map or {})
    def header_candidates(accepted: set[str]) -> list[str]:
        return [str(column) for column in raw.columns if _fold(str(column)) in accepted]
    canonical_name_headers = header_candidates({"segmentname", "分段名称", "标签名称"})
    legacy_name_headers = header_candidates({"name", "名称", "segmentlabel"})
    if "segment_name" not in requested:
        candidates = canonical_name_headers or legacy_name_headers
        if len(candidates) > 1:
            raise DataImportError("More than one source column can map to segment_name; map it explicitly")
        if candidates:
            requested["segment_name"] = candidates[0]
    elif requested["segment_name"] not in raw.columns:
        raise DataImportError("Mapped segment_name source column does not exist")
    name_alias_source = requested.get("name")
    if name_alias_source is None and legacy_name_headers:
        if len(legacy_name_headers) > 1:
            raise DataImportError("More than one legacy name column is ambiguous")
        name_alias_source = legacy_name_headers[0]
    required = ["start_s", "end_s"]
    if scope_mode == "participant_session":
        if subject_id is None:
            required.append("subject_id")
        if session_id is None:
            required.append("session_id")
    else:
        required.append("interaction_id")
    optional = ("segment_id", "segment_name", "subject_id", "session_id", "interaction_id",
                "condition", "trial", "stimulus_id", "baseline_start_s", "baseline_end_s",
                "corrected_from_segment_id", "timebase_id", "segment_kind", "physical_semantics_ref",
                "overlap_policy")
    mapping = _resolve_mapping(raw.columns, requested, tuple(required), optional)
    if "segment_name" not in mapping:
        raise DataImportError("New segment imports require segment_name (or the legacy name alias)")
    out = pd.DataFrame(index=raw.index)
    for name in ("segment_id", "segment_name", "start_s", "end_s", "subject_id", "session_id", "interaction_id",
                 "condition", "trial", "stimulus_id", "baseline_start_s", "baseline_end_s",
                 "corrected_from_segment_id", "timebase_id", "segment_kind", "physical_semantics_ref",
                 "overlap_policy"):
        source = mapping.get(name)
        if source is not None and name not in {"start_s", "end_s"}:
            out[name] = raw[source]
    if name_alias_source is not None:
        if name_alias_source not in raw.columns:
            raise DataImportError("Mapped legacy name source column does not exist")
        out["name"] = raw[name_alias_source]
    elif "segment_name" in out:
        out["name"] = out["segment_name"]
    if scope_mode == "participant_session":
        for name, explicit in (("subject_id", subject_id), ("session_id", session_id)):
            if name not in out:
                if explicit is None:
                    raise DataImportError(f"{name} is absent from segments; explicit assignment is required")
                out[name] = str(explicit)
            else:
                blanks = out[name].isna() | out[name].astype("string").str.strip().eq("")
                if blanks.any():
                    if explicit is None:
                        raise DataImportError(f"Blank {name} in segments; no ID will be inferred")
                    out.loc[blanks, name] = str(explicit)
    else:
        for name in ("subject_id", "session_id"):
            if name not in out:
                out[name] = pd.NA
    for name in ("interaction_id", "condition", "trial", "stimulus_id"):
        if name in mapping:
            out[name] = raw[mapping[name]]
        elif name == "interaction_id":
            continue
        else:
            out[name] = pd.NA
    if time.time_unit == "unknown":
        raise DataImportError("Confirm time_unit for segment boundaries")
    out["start_s"], start_err = _convert_time(raw[mapping["start_s"]], time)
    out["end_s"], end_err = _convert_time(raw[mapping["end_s"]], time)
    if start_err.any() or end_err.any():
        raise DataImportError("Segment boundaries contain invalid times")
    if (out["end_s"] <= out["start_s"]).any():
        raise DataImportError("Segments must be nonempty left-closed, right-open intervals")
    for name in ("baseline_start_s", "baseline_end_s"):
        if name in mapping:
            out[name], bad = _convert_time(raw[mapping[name]], time)
            if bad.any():
                raise DataImportError(f"{name} contains invalid times")
        elif name not in out:
            out[name] = pd.NA
    if (out["baseline_start_s"].isna() ^ out["baseline_end_s"].isna()).any():
        raise DataImportError("Label baseline windows require both baseline_start_s and baseline_end_s")
    if (out["baseline_end_s"].dropna() <= out.loc[out["baseline_end_s"].notna(), "baseline_start_s"]).any():
        raise DataImportError("Label baseline windows must be nonempty")
    audit_fields = {
        "source_sha256": imported.sha256,
        "source_file_name": imported.source_name,
        "source_sheet": imported.sheet_name,
        "source_row": list(imported.source_rows or range(len(raw))),
        "raw_start": raw[mapping["start_s"]].to_list(),
        "raw_end": raw[mapping["end_s"]].to_list(),
        "raw_segment_name": raw[mapping["segment_name"]].to_list(),
    }
    for column, value in audit_fields.items():
        out[column] = value
    if timebase_id is None:
        timebase_id = time.clock_source or time.frame_id
    try:
        normalized, scope_audit = normalize_segment_records(
            out,
            scope_config=scope_config,
            member_session_map=member_session_map,
            dataset_scope=dataset_scope,
            timebase_id=timebase_id,
            source_audit={"sha256": imported.sha256, "source_name": imported.source_name,
                          "sheet_name": imported.sheet_name},
        )
    except ScopeError as exc:
        raise DataImportError(str(exc)) from exc
    normalized.attrs["segment_scope_audit"] = scope_audit
    normalized.attrs["schema_version"] = "1.2.0"
    normalized.attrs["source_manifest"] = imported.manifest_entry()
    return normalized.reset_index(drop=True)


def normalize_aois(
    imported: ImportedTable,
    *,
    column_map: Mapping[str, str] | None = None,
    coordinate_unit: str = "unknown",
    coordinate_frame: str = "unknown",
    overlap_policy: str | None = None,
    missing_name_policy: Literal["require_source", "allow_legacy_missing"] = "require_source",
) -> pd.DataFrame:
    if imported.role != "aois":
        raise DataImportError("normalize_aois requires an AOI table")
    raw = imported.frame
    if raw.empty:
        source = _aoi_source_context(imported, {}, None)
        raise AOIGeometryImportError([AOIGeometryIssue(
            "aoi_table_empty", "AOI source table has no data rows",
            source=source, canonical_field=None,
        )])
    if len(raw) > 1024:
        source = _aoi_source_context(imported, {}, None)
        raise AOIGeometryImportError([AOIGeometryIssue(
            "aoi_geometry_resource_limit", "AOI row limit exceeded; rows were not truncated",
            source=source, canonical_field=None,
            details={"limit_name": "max_aoi_rows", "limit": 1024, "observed": len(raw)},
        )])
    try:
        required_file_fields = AOI_REQUIRED_COLUMNS + (("aoi_name",) if missing_name_policy == "require_source" else ())
        mapping = _resolve_mapping(
            raw.columns, column_map or {}, required_file_fields,
            ("xmin", "ymin", "xmax", "ymax", "vertices", "priority", "overlap_policy", "aoi_name"),
        )
    except DataImportError as exc:
        source = _aoi_source_context(imported, {}, None)
        issue = AOIGeometryIssue(
            "aoi_source_mapping_invalid", str(exc), source=source,
            canonical_field=None,
            details={"mapping_error": str(exc)},
        )
        raise AOIGeometryImportError([issue]) from exc
    if overlap_policy not in {None, "exclusive", "multi_label", "unresolved"}:
        source = _aoi_source_context(imported, mapping, None)
        issue = AOIGeometryIssue(
            "aoi_overlap_policy_invalid", "overlap_policy must be exclusive, multi_label, or unresolved",
            source=source, canonical_field="overlap_policy",
        )
        raise AOIGeometryImportError([issue])
    out = pd.DataFrame(index=raw.index)
    for name in AOI_REQUIRED_COLUMNS:
        out[name] = raw[mapping[name]]
    out["aoi_name"] = raw[mapping["aoi_name"]] if "aoi_name" in mapping else pd.NA
    for name in ("xmin", "ymin", "xmax", "ymax", "vertices", "priority"):
        out[name] = raw[mapping[name]] if name in mapping else pd.NA
    if "overlap_policy" in mapping:
        out["overlap_policy"] = raw[mapping["overlap_policy"]]
    else:
        out["overlap_policy"] = "unresolved"
    if overlap_policy is not None:
        out["overlap_policy"] = overlap_policy
    out["unit"] = coordinate_unit
    out["coordinate_frame"] = coordinate_frame
    rows: list[dict[str, Any]] = []
    for position, (_, source_row) in enumerate(out.iterrows()):
        row = source_row.to_dict()
        if str(row.get("shape", "")).strip().lower() == "rectangle":
            for field_name in ("xmin", "ymin", "xmax", "ymax"):
                value = row.get(field_name)
                row[field_name] = _aoi_decimal_cell(
                    value, field_name=field_name,
                    source=_aoi_source_context(imported, mapping, position),
                    source_column=mapping.get(field_name),
                )
        priority_raw = row.get("priority")
        if not _aoi_cell_missing(priority_raw):
            try:
                row["priority"] = _aoi_decimal_cell(
                    priority_raw, field_name="priority",
                    source=_aoi_source_context(imported, mapping, position),
                    source_column=mapping.get("priority"),
                )
            except AOIGeometryImportError:
                # Priority is required only when an exclusive overlap is resolved.
                # Keep invalid raw priority unavailable so the labeling policy emits
                # a scoped, reasoned overlap error instead of guessing.
                row["priority"] = pd.NA
        else:
            row["priority"] = pd.NA
        row["aoi_source"] = _aoi_source_context(imported, mapping, position)
        rows.append(row)
    source_context = {
        "source_name": imported.source_name,
        "source_file": str(imported.source_path),
        "sha256": imported.sha256,
        "sheet": imported.sheet_name,
        "source_columns": dict(mapping),
        "origin": "ImportedTable",
    }
    try:
        prepared = prepare_static_aois(
            rows, source=source_context, missing_name_policy=missing_name_policy,
        )
    except AOITableGeometryError as exc:
        raise AOIGeometryImportError(exc.issues) from exc
    except AOIGeometryError as exc:
        # Topology/resource failures can be raised immediately before the
        # table-wide issue aggregation path. Keep the public IO contract
        # uniform so callers always receive a DataImportError with provenance.
        raise AOIGeometryImportError([exc.issue]) from exc
    normalized = pd.DataFrame([dict(row) for row in prepared.rows]).drop(columns=["_canonical_geometry"], errors="ignore")
    normalized = normalized.reset_index(drop=True)
    normalized.attrs["aoi_domain_schema_version"] = AOI_DOMAIN_SCHEMA_VERSION
    normalized.attrs["aoi_geometry_version"] = AOI_GEOMETRY_VERSION
    normalized.attrs["aoi_geometry_audit"] = dict(prepared.audit)
    normalized.attrs["source_manifest"] = imported.manifest_entry()
    return normalized


def normalize_performance(imported: ImportedTable) -> pd.DataFrame:
    if imported.role != "performance":
        raise DataImportError("normalize_performance requires a performance table")
    frame = imported.frame.copy()
    if frame.columns.has_duplicates:
        raise DataImportError("Duplicate performance headers are ambiguous")
    if frame.empty:
        raise DataImportError("Performance table is empty")
    return frame


def _read_csv(path: Path, options: CsvOptions, preview_rows: int | None) -> pd.DataFrame:
    try:
        kwargs: dict[str, Any] = {
            "encoding": options.encoding,
            "sep": options.delimiter,
            "decimal": options.decimal,
            "dtype": object,
            "keep_default_na": False,
            "skip_blank_lines": False,
            "engine": "python" if options.delimiter is None else "c",
        }
        if preview_rows is not None:
            if preview_rows < 1:
                raise DataImportError("preview_rows must be positive")
            kwargs["nrows"] = preview_rows
        frame = pd.read_csv(path, **kwargs)
    except (UnicodeError, ValueError, pd.errors.ParserError) as exc:
        raise DataImportError(f"Could not parse CSV with the selected options: {exc}") from exc
    frame.columns = [str(name).strip() for name in frame.columns]
    return frame


def _read_xlsx(path: Path, sheet_name: str | None, preview_rows: int | None) -> tuple[pd.DataFrame, str, tuple[int, ...]]:
    _guard_xlsx_container(path)
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise DataImportError(f"Could not open XLSX safely: {exc}") from exc
    try:
        names = workbook.sheetnames
        if not names:
            raise DataImportError("XLSX contains no worksheets")
        if sheet_name is None:
            if len(names) != 1:
                raise DataImportError("Select a worksheet explicitly when the workbook has multiple sheets")
            sheet_name = names[0]
        if sheet_name not in names:
            raise DataImportError(f"Worksheet not found: {sheet_name}")
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows()
        first = next(iterator, None)
        if first is None:
            return pd.DataFrame(), sheet_name, ()
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in first]
        if any(not name for name in headers):
            raise DataImportError("XLSX header row contains blank column names")
        if len(set(headers)) != len(headers):
            raise DataImportError("XLSX contains duplicate headers")
        records: list[list[Any]] = []
        source_rows: list[int] = []
        max_rows = preview_rows
        if max_rows is not None and max_rows < 1:
            raise DataImportError("preview_rows must be positive")
        for row_number, row in enumerate(iterator, start=2):
            if max_rows is not None and len(records) >= max_rows:
                break
            values: list[Any] = []
            for cell in row[:len(headers)]:
                if cell.data_type == "f":
                    raise DataImportError(
                        f"Formula cell {cell.coordinate} is not evaluated; replace with an explicit value"
                    )
                values.append(cell.value)
            values.extend([None] * (len(headers) - len(values)))
            records.append(values)
            source_rows.append(row_number)
        frame = pd.DataFrame.from_records(records, columns=headers)
        return frame, sheet_name, tuple(source_rows)
    finally:
        workbook.close()


def _guard_xlsx_container(path: Path) -> None:
    """Reject ZIP bombs and active content before handing XLSX to openpyxl."""
    try:
        with zipfile.ZipFile(path, "r") as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_MEMBERS:
                raise DataImportError("XLSX contains too many ZIP members")
            total_size = 0
            names: set[str] = set()
            for entry in entries:
                name = entry.filename
                folded_name = name.casefold()
                member_posix = PurePosixPath(name)
                member_windows = PureWindowsPath(name)
                if (
                    "\\" in name
                    or member_posix.is_absolute()
                    or member_windows.is_absolute()
                    or member_windows.drive
                    or ".." in member_posix.parts
                    or ".." in member_windows.parts
                ):
                    raise DataImportError("XLSX contains an unsafe ZIP member path")
                if folded_name in names:
                    raise DataImportError("XLSX contains duplicate ZIP member names")
                names.add(folded_name)
                if entry.flag_bits & 0x1:
                    raise DataImportError("Encrypted XLSX containers are not supported")
                if entry.file_size < 0 or entry.compress_size < 0:
                    raise DataImportError("XLSX contains invalid ZIP size metadata")
                if entry.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
                    raise DataImportError("XLSX uses an unsupported ZIP compression method")
                total_size += entry.file_size
                if total_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise DataImportError("XLSX expands beyond the bounded uncompressed-size limit")
                if entry.file_size and entry.compress_size == 0:
                    raise DataImportError("XLSX contains an invalid zero-size compressed member")
                if entry.compress_size and entry.file_size / entry.compress_size > MAX_XLSX_COMPRESSION_RATIO:
                    raise DataImportError("XLSX contains a suspiciously high compression-ratio member")
                if folded_name.startswith(("xl/activex/", "xl/embeddings/")) or folded_name == "xl/vbaproject.bin":
                    raise DataImportError("XLSX active content and embedded objects are not supported")
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise DataImportError("File is not a valid XLSX workbook container")
    except DataImportError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise DataImportError(f"Could not inspect XLSX ZIP container: {exc}") from exc


def _resolve_mapping(
    columns: Any,
    requested: Mapping[str, str],
    required: tuple[str, ...],
    optional: tuple[str, ...],
) -> dict[str, str]:
    existing = [str(column) for column in columns]
    if len(set(existing)) != len(existing):
        raise DataImportError("Duplicate source headers are ambiguous")
    normalized: dict[str, list[str]] = {}
    for column in existing:
        normalized.setdefault(_fold(column), []).append(column)
    mapping: dict[str, str] = {}
    aliases = {**COLUMN_ALIASES}
    aliases.setdefault("start_s", ("start_s", "start_time", "开始时间"))
    aliases.setdefault("end_s", ("end_s", "end_time", "结束时间"))
    aliases.setdefault("segment_id", ("segment_id", "segment", "分段编号"))
    aliases.setdefault("aoi_id", ("aoi_id", "aoi", "roi_id", "兴趣区编号"))
    aliases.setdefault("aoi_name", ("aoi_name", "name", "label", "名称", "兴趣区名称", "区域名称"))
    aliases.setdefault("overlap_policy", ("overlap_policy", "重叠策略"))
    aliases.setdefault("shape", ("shape", "形状"))
    aliases.setdefault("xmin", ("xmin", "x_min", "left"))
    aliases.setdefault("ymin", ("ymin", "y_min", "top"))
    aliases.setdefault("xmax", ("xmax", "x_max", "right"))
    aliases.setdefault("ymax", ("ymax", "y_max", "bottom"))
    aliases.setdefault("vertices", ("vertices", "points", "顶点"))
    aliases.setdefault("priority", ("priority", "优先级"))
    for canonical in (*required, *optional):
        if canonical in requested:
            source = requested[canonical]
            if source not in existing:
                raise DataImportError(f"Mapped source column {source!r} for {canonical!r} does not exist")
            mapping[canonical] = source
            continue
        candidates = []
        for alias in (*aliases.get(canonical, ()), canonical):
            candidates.extend(normalized.get(_fold(alias), []))
        candidates = list(dict.fromkeys(candidates))
        if len(candidates) > 1:
            raise DataImportError(f"More than one source column can map to {canonical!r}; map it explicitly")
        if candidates:
            mapping[canonical] = candidates[0]
        elif canonical in required:
            raise DataImportError(f"Required field {canonical!r} is missing; map the source column explicitly")
    reverse: dict[str, str] = {}
    for canonical, source in mapping.items():
        previous = reverse.get(source)
        if previous is not None and previous != canonical:
            raise DataImportError(
                f"Source column {source!r} is mapped to both {previous!r} and {canonical!r}"
            )
        reverse[source] = canonical
    return mapping


def _aoi_source_context(
    imported: ImportedTable,
    mapping: Mapping[str, str],
    position: int | None,
) -> dict[str, Any]:
    source_row = None
    if position is not None and 0 <= position < len(imported.source_rows):
        source_row = imported.source_rows[position]
    return {
        "source_name": imported.source_name,
        "source_file": str(imported.source_path),
        "sha256": imported.sha256,
        "sheet": imported.sheet_name,
        "source_row": source_row,
        "source_columns": dict(mapping),
        "origin": "ImportedTable",
    }


def _aoi_cell_missing(value: Any) -> bool:
    if value is None or value is pd.NA:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        missing = pd.isna(value)
        return bool(missing) if isinstance(missing, (bool, np.bool_)) else False
    except (TypeError, ValueError):
        return False


def _aoi_decimal_cell(
    value: Any,
    *,
    field_name: str,
    source: Mapping[str, Any],
    source_column: str | None,
) -> float:
    """Convert a finite scalar or conservative decimal-text form to binary64."""
    import math

    from .aoi_geometry import AOIGeometryIssue

    message = "AOI bounds must be finite numeric values; expressions and booleans are invalid"
    if isinstance(value, (bool, np.bool_)) or _aoi_cell_missing(value):
        issue = AOIGeometryIssue(
            "aoi_geometry_coordinate_invalid", message, source=source,
            canonical_field=field_name, source_column=source_column,
        )
        raise AOIGeometryImportError([issue])
    if isinstance(value, str):
        text = value.strip()
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", text):
            issue = AOIGeometryIssue(
                "aoi_geometry_coordinate_invalid", message, source=source,
                canonical_field=field_name, source_column=source_column,
            )
            raise AOIGeometryImportError([issue])
        try:
            converted = float(text)
        except ValueError as exc:
            issue = AOIGeometryIssue(
                "aoi_geometry_coordinate_invalid", message, source=source,
                canonical_field=field_name, source_column=source_column,
            )
            raise AOIGeometryImportError([issue]) from exc
    elif isinstance(value, (int, float, np.integer, np.floating)):
        try:
            converted = float(value)
        except (TypeError, ValueError, OverflowError) as exc:
            issue = AOIGeometryIssue(
                "aoi_geometry_coordinate_invalid", message, source=source,
                canonical_field=field_name, source_column=source_column,
            )
            raise AOIGeometryImportError([issue]) from exc
    else:
        issue = AOIGeometryIssue(
            "aoi_geometry_coordinate_invalid", message, source=source,
            canonical_field=field_name, source_column=source_column,
        )
        raise AOIGeometryImportError([issue])
    if not math.isfinite(converted):
        issue = AOIGeometryIssue(
            "aoi_geometry_coordinate_nonfinite", "AOI coordinates must be finite binary64 values",
            source=source, canonical_field=field_name, source_column=source_column,
        )
        raise AOIGeometryImportError([issue])
    return 0.0 if converted == 0.0 else converted


def _convert_time(values: pd.Series, metadata: CoordinateMetadata) -> tuple[pd.Series, pd.Series]:
    raw = values.copy()
    if metadata.time_unit in {"s", "ms", "frame"}:
        numeric = pd.to_numeric(raw, errors="coerce")
        if metadata.time_unit == "ms":
            numeric = numeric / 1000.0
        elif metadata.time_unit == "frame":
            # Constructor validation guarantees a positive, explicit frame rate.
            numeric = numeric / float(metadata.frame_rate_hz)
        error = raw.notna() & raw.astype("string").str.strip().ne("") & numeric.isna()
        numeric, nonfinite = _finite_numeric_series(numeric)
        return numeric, (error | nonfinite).astype("boolean")
    if metadata.time_unit != "datetime":
        raise DataImportError("Unsupported time unit")
    if not metadata.time_origin:
        raise DataImportError("datetime timestamps require an explicit time_origin")
    try:
        parsed = pd.to_datetime(raw, errors="coerce")
        origin = pd.Timestamp(metadata.time_origin)
        if parsed.dt.tz is None:
            if not metadata.timezone:
                raise DataImportError("Naive datetime requires an explicit timezone")
            parsed = parsed.dt.tz_localize(metadata.timezone, ambiguous="NaT", nonexistent="NaT")
        else:
            parsed = parsed.dt.tz_convert("UTC")
        if origin.tzinfo is None:
            if not metadata.timezone:
                raise DataImportError("Naive time_origin requires an explicit timezone")
            origin = origin.tz_localize(metadata.timezone)
        origin = origin.tz_convert("UTC")
        seconds = (parsed.dt.tz_convert("UTC") - origin).dt.total_seconds()
        error = raw.notna() & raw.astype("string").str.strip().ne("") & seconds.isna()
        seconds, nonfinite = _finite_numeric_series(seconds)
        return seconds, (error | nonfinite).astype("boolean")
    except (TypeError, ValueError, AttributeError) as exc:
        if isinstance(exc, DataImportError):
            raise
        raise DataImportError(f"Invalid datetime or time_origin: {exc}") from exc


def _finite_numeric_series(values: pd.Series) -> tuple[pd.Series, pd.Series]:
    """Keep parsed non-finite numbers out of analysis tables while retaining a row flag."""
    numeric = pd.to_numeric(values, errors="coerce")
    finite = pd.Series(
        np.isfinite(numeric.to_numpy(dtype=float, na_value=np.nan)),
        index=numeric.index,
        dtype="boolean",
    )
    nonfinite = numeric.notna() & ~finite
    return numeric.mask(nonfinite).astype("Float64"), nonfinite.astype("boolean")


def _decode_validity(values: pd.Series, custom: Mapping[str, bool]) -> pd.Series:
    defaults = {"true": True, "false": False, "1": True, "0": False,
                "yes": True, "no": False, "valid": True, "invalid": False}
    mapping = {**defaults, **{str(key).strip().casefold(): value for key, value in custom.items()}}
    decoded: list[Any] = []
    bad: list[str] = []
    for value in values:
        if pd.isna(value) or (isinstance(value, str) and not value.strip()):
            decoded.append(pd.NA)
            continue
        if isinstance(value, bool):
            decoded.append(value)
            continue
        key = str(value).strip().casefold()
        if key in mapping:
            decoded.append(bool(mapping[key]))
        elif isinstance(value, (int, float)) and value in (0, 1):
            decoded.append(bool(value))
        else:
            bad.append(key)
            decoded.append(pd.NA)
    if bad:
        raise DataImportError(
            "Validity codes are ambiguous; provide an explicit validity_values mapping for: "
            + ", ".join(sorted(set(bad))[:8])
        )
    return pd.Series(pd.array(decoded, dtype="boolean"), index=values.index)


def _decode_reading_direction(values: pd.Series, custom: Mapping[str, str]) -> pd.Series:
    aliases = {str(key).strip().casefold(): str(value).strip().casefold() for key, value in custom.items()}
    decoded: list[Any] = []
    unknown: list[str] = []
    for value in values:
        if pd.isna(value) or (isinstance(value, str) and not value.strip()):
            decoded.append(pd.NA)
            continue
        key = str(value).strip().casefold()
        direction = aliases.get(key, key)
        if direction not in READING_DIRECTIONS:
            unknown.append(key)
            decoded.append(pd.NA)
        else:
            decoded.append(direction)
    if unknown:
        raise DataImportError(
            "Reading-direction codes are ambiguous; provide reading_direction_values for: "
            + ", ".join(sorted(set(unknown))[:8])
        )
    return pd.Series(pd.array(decoded, dtype="string"), index=values.index)


def _fold(value: str) -> str:
    return re.sub(r"[\s\-_./]+", "", str(value)).casefold()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()
