"""Bounded, provenance-aware exports for immutable GUI run records.

The functions in this module serialize values already present in a RunRecord.
They do not run analysis methods or read ResultPage's sorted/filter proxy model.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
import base64
import csv
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile
from typing import Any, Callable, Literal, Mapping, Sequence

import numpy as np
import pandas as pd

from pyeyedata.gui.models import RunRecord


ExportFormat = Literal["csv", "xlsx"]
SelectionScope = Literal["full", "filtered", "selected"]
XlsxOverflow = Literal["reject", "split_sheets"]
CancelCallback = Callable[[], bool]
ProgressCallback = Callable[[float, str], None]

XLSX_MAX_ROWS = 1_048_576  # Includes the header row.
XLSX_MAX_COLUMNS = 16_384
XLSX_MAX_CELL_TEXT = 32_767
_XLSX_DATA_ROWS = XLSX_MAX_ROWS - 1
_LONG_TEXT_CHUNK = 32_000

_IDENTITY_COLUMNS = (
    "run_id", "metric_id", "subject_id", "session_id", "interaction_id", "pair_id",
    "segment_id", "segment_instance_id", "trial", "condition", "analysis_stream_id",
    "eye_label", "signal", "event_id", "observation_id",
)
_ID_COLUMNS = frozenset(_IDENTITY_COLUMNS) | frozenset({
    "participant_id", "member_id", "label_id", "source_id", "source_event_id",
    "physical_partition_id", "family_id", "reference_id", "aoi_id",
})
_REVIEW_COLUMNS = (
    "run_id", "revision_id", "metric_id", "subject_id", "session_id", "interaction_id",
    "pair_id", "segment_id", "segment_instance_id", "trial", "condition", "label",
    "analysis_stream_id", "eye_label", "signal", "event_id", "observation_id", "value",
    "metric_value", "unit", "status", "availability_status", "reason_code", "reason",
    "n", "n_raw", "n_interpolated", "n_observed", "n_effective", "independent_n",
    "effective_n", "effective_duration_s", "duration_s", "algorithm", "algorithm_id",
    "parameter_ref", "source_ref", "family", "family_id", "estimate", "ci_lower",
    "ci_upper", "p_value", "p_adjusted", "p_raw", "p_adjusted_value",
)


class ExportError(ValueError):
    """A request cannot be faithfully or safely exported."""


class ExportIdentityError(ExportError):
    """A request no longer identifies the RunRecord/table that created it."""


class ExportLimitError(ExportError):
    """An output limit would be exceeded without an explicit representation."""


class ExportCancelled(InterruptedError):
    """Cooperative cancellation before output promotion; prior files are intact."""


def _export_checkpoint(cancel: CancelCallback | None = None,
                       progress: ProgressCallback | None = None,
                       fraction: float = 0.0, message: str = "") -> None:
    if cancel is not None and cancel():
        raise ExportCancelled("export cancelled before commit")
    if progress is not None:
        progress(fraction, message)
    if cancel is not None and cancel():
        raise ExportCancelled("export cancelled before commit")


@dataclass(frozen=True, slots=True)
class TableExportRequest:
    """Immutable export dispatch captured from one RunRecord and table view."""

    run_id: str
    revision_id: str
    config_fingerprint: str
    table_name: str
    file_format: ExportFormat
    scope: SelectionScope
    source_positions: tuple[int, ...]
    source_row_ids: tuple[str, ...]
    table_fingerprint: str
    columns: tuple[str, ...] | None
    reference_tables: tuple[str, ...]
    filter_description: str | None
    destination: str
    overwrite: bool = False
    xlsx_overflow: XlsxOverflow = "reject"


@dataclass(frozen=True, slots=True)
class ExportPreview:
    run_id: str
    revision_id: str
    table_name: str
    file_format: ExportFormat
    scope: SelectionScope
    row_count: int
    source_row_count: int
    requested_columns: tuple[str, ...]
    output_columns: tuple[str, ...]
    automatically_included_review_columns: tuple[str, ...]
    reference_tables: tuple[str, ...]
    worksheet_count: int
    split_required: bool
    over_limit_cells: int
    formula_like_text_cells: int
    warnings: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ExportResult:
    path: Path
    audit_path: Path
    references_path: Path | None
    sha256: str
    rows_written: int
    columns_written: tuple[str, ...]
    formula_like_text_cells: int
    long_text_cells: int
    worksheet_names: tuple[str, ...]

    @property
    def commit_complete(self) -> bool:
        """This receipt is returned only after the entire output group commits."""
        return True


def _missing(value: Any) -> bool:
    if value is None or value is pd.NA or value is pd.NaT:
        return True
    try:
        result = pd.isna(value)
        return bool(result) if isinstance(result, (bool, np.bool_)) else False
    except (TypeError, ValueError):
        return False


def _safe_value(value: Any) -> Any:
    """Typed, JSON-safe form used only for identities and audit records."""
    if _missing(value):
        return {"$missing": type(value).__name__}
    if isinstance(value, np.generic):
        return _safe_value(value.item())
    if isinstance(value, pd.Timestamp):
        return {"$datetime": value.isoformat()}
    if isinstance(value, (datetime, date)):
        return {"$datetime": value.isoformat()}
    if isinstance(value, Decimal):
        return {"$decimal": str(value)}
    if isinstance(value, bytes):
        return {"$bytes": base64.b64encode(value).decode("ascii")}
    if isinstance(value, bytearray):
        return {"$bytes": base64.b64encode(bytes(value)).decode("ascii")}
    if isinstance(value, float) and not math.isfinite(value):
        return {"$float": repr(value)}
    if isinstance(value, Mapping):
        return {str(key): _safe_value(item) for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))}
    if isinstance(value, (list, tuple)):
        return [_safe_value(item) for item in value]
    if isinstance(value, (set, frozenset)):
        items = [_safe_value(item) for item in value]
        return sorted(items, key=lambda item: json.dumps(item, sort_keys=True, ensure_ascii=False))
    if isinstance(value, (str, bool, int, float)):
        return value
    return {"$repr": str(value), "$type": type(value).__qualname__}


def _identity_value(value: Any) -> Any:
    """Match ResultsPage's row-identity serialization for proxy/source binding."""
    if _missing(value):
        return None
    if isinstance(value, np.generic):
        return _identity_value(value.item())
    if isinstance(value, bytes):
        return {"type": "bytes", "length": len(value), "sha256": hashlib.sha256(value).hexdigest()}
    if isinstance(value, bytearray):
        raw = bytes(value)
        return {"type": "bytes", "length": len(raw), "sha256": hashlib.sha256(raw).hexdigest()}
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return str(value)
    if isinstance(value, Mapping):
        return {str(key): _identity_value(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set, frozenset)):
        return [_identity_value(item) for item in value]
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "to_dict") and callable(value.to_dict):
        try:
            return _identity_value(value.to_dict())
        except (TypeError, ValueError):
            pass
    return str(value)


def source_row_identity(frame: pd.DataFrame, row_position: int) -> str:
    """Return the stable source-table row key used by ResultsPage selections."""
    if not 0 <= int(row_position) < len(frame):
        raise IndexError(f"source row {row_position} is outside the table")
    row = frame.iloc[int(row_position)]
    identity_columns = [name for name in _IDENTITY_COLUMNS
                        if name in frame.columns and not _missing(row[name])]
    if identity_columns:
        identity = {name: _identity_value(row[name]) for name in identity_columns}
    else:
        identity = {str(name): _identity_value(value) for name, value in row.items()}
    packed = json.dumps(identity, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(packed.encode("utf-8")).hexdigest()


def _table_fingerprint(frame: pd.DataFrame, *, cancel: CancelCallback | None = None) -> str:
    digest = hashlib.sha256()
    header = {
        "columns": [str(item) for item in frame.columns],
        "dtypes": [str(dtype) for dtype in frame.dtypes],
        "index_dtype": str(frame.index.dtype),
        "rows": int(len(frame)),
    }
    digest.update(json.dumps(header, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode("utf-8"))
    digest.update(b"\n")
    for index, row in zip(frame.index, frame.itertuples(index=False, name=None)):
        _export_checkpoint(cancel)
        packed = json.dumps([_safe_value(index), *(_safe_value(value) for value in row)],
                            ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)
        digest.update(packed.encode("utf-8"))
        digest.update(b"\n")
    return digest.hexdigest()


def create_table_export_request(
    record: RunRecord,
    table_name: str,
    destination: str | Path,
    *,
    file_format: ExportFormat = "csv",
    scope: SelectionScope = "full",
    source_positions: Sequence[int] | None = None,
    source_row_ids: Sequence[str] | None = None,
    columns: Sequence[str] | None = None,
    reference_tables: Sequence[str] = (),
    filter_description: str | None = None,
    overwrite: bool = False,
    xlsx_overflow: XlsxOverflow = "reject",
) -> TableExportRequest:
    """Capture table identity and row mapping before asynchronous export dispatch."""
    if file_format not in ("csv", "xlsx"):
        raise ExportError("file_format must be 'csv' or 'xlsx'")
    if scope not in ("full", "filtered", "selected"):
        raise ExportError("scope must be 'full', 'filtered' or 'selected'")
    if xlsx_overflow not in ("reject", "split_sheets"):
        raise ExportError("xlsx_overflow must be 'reject' or 'split_sheets'")
    names = tuple(record.table_names())
    if table_name not in names:
        raise ExportIdentityError(f"run {record.run_id!r} has no table named {table_name!r}")
    frame = record.table_copy(table_name)
    if not isinstance(frame, pd.DataFrame):
        raise ExportError(f"record table {table_name!r} is not a DataFrame")
    string_columns = tuple(str(column) for column in frame.columns)
    if len(set(string_columns)) != len(string_columns):
        raise ExportError("table has duplicate or string-colliding column names; export would make the schema ambiguous")
    positions = tuple(int(value) for value in (source_positions or ()))
    supplied_ids = tuple(str(value) for value in (source_row_ids or ()))
    if scope == "full":
        if positions or supplied_ids:
            raise ExportError("full-table requests do not accept a row subset")
    else:
        if len(set(positions)) != len(positions):
            raise ExportError("source_positions must not contain duplicates")
        if any(position < 0 or position >= len(frame) for position in positions):
            raise ExportIdentityError("a selected source row is outside the immutable table")
        resolved_ids = tuple(source_row_identity(frame, position) for position in positions)
        if supplied_ids and supplied_ids != resolved_ids:
            raise ExportIdentityError("captured row IDs do not match the selected source positions")
        supplied_ids = resolved_ids
    requested_columns = None if columns is None else tuple(str(column) for column in columns)
    if requested_columns is not None:
        available = {str(column) for column in frame.columns}
        missing_columns = [column for column in requested_columns if column not in available]
        if missing_columns:
            raise ExportError(f"unknown selected columns: {', '.join(missing_columns)}")
        if len(set(requested_columns)) != len(requested_columns):
            raise ExportError("columns must not contain duplicates")
    refs = tuple(str(name) for name in reference_tables)
    if len(set(refs)) != len(refs):
        raise ExportError("reference_tables must not contain duplicates")
    missing_refs = [name for name in refs if name not in names]
    if missing_refs:
        raise ExportIdentityError(f"missing named reference tables: {', '.join(missing_refs)}")
    if table_name in refs:
        raise ExportError("the main table cannot also be a linked reference table")
    for reference_name in refs:
        reference_frame = record.table_copy(reference_name)
        reference_columns = tuple(str(column) for column in reference_frame.columns)
        if len(set(reference_columns)) != len(reference_columns):
            raise ExportError(
                f"reference table {reference_name!r} has duplicate or string-colliding column names; "
                "export would make the schema ambiguous"
            )
    dest = Path(destination).expanduser().absolute()
    return TableExportRequest(
        run_id=str(record.run_id), revision_id=str(record.revision_id),
        config_fingerprint=str(record.config_fingerprint), table_name=str(table_name),
        file_format=file_format, scope=scope, source_positions=positions,
        source_row_ids=supplied_ids, table_fingerprint=_table_fingerprint(frame),
        columns=requested_columns, reference_tables=refs,
        filter_description=filter_description, destination=str(dest),
        overwrite=bool(overwrite), xlsx_overflow=xlsx_overflow,
    )


def _verify_request(record: RunRecord, request: TableExportRequest, *,
                    cancel: CancelCallback | None = None) -> pd.DataFrame:
    _export_checkpoint(cancel)
    if (str(record.run_id), str(record.revision_id), str(record.config_fingerprint)) != (
        request.run_id, request.revision_id, request.config_fingerprint
    ):
        raise ExportIdentityError("export request belongs to a different run revision or configuration")
    if request.table_name not in record.table_names():
        raise ExportIdentityError(f"run {record.run_id!r} no longer contains {request.table_name!r}")
    frame = record.table_copy(request.table_name)
    if _table_fingerprint(frame, cancel=cancel) != request.table_fingerprint:
        raise ExportIdentityError("source table changed after export selection was captured")
    if request.scope != "full":
        if len(request.source_positions) != len(request.source_row_ids):
            raise ExportIdentityError("captured source positions and row identities have different lengths")
        current = []
        for position in request.source_positions:
            _export_checkpoint(cancel)
            current.append(source_row_identity(frame, position))
        if tuple(current) != request.source_row_ids:
            raise ExportIdentityError("selected row identities no longer match the immutable source table")
    return frame


def _output_columns(frame: pd.DataFrame, requested: tuple[str, ...] | None) -> tuple[tuple[str, ...], tuple[str, ...]]:
    available = [str(column) for column in frame.columns]
    selected = list(available if requested is None else requested)
    selected_set = set(selected)
    review = [name for name in _REVIEW_COLUMNS if name in available and name not in selected_set]
    return tuple(selected + review), tuple(review)


def _selection_positions(request: TableExportRequest, frame: pd.DataFrame) -> Sequence[int] | range:
    return range(len(frame)) if request.scope == "full" else request.source_positions


def _to_text(value: Any) -> str:
    if _missing(value):
        return ""
    if isinstance(value, np.generic):
        return _to_text(value.item())
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isinf(value):
            return "Infinity" if value > 0 else "-Infinity"
        return repr(value)
    if isinstance(value, Mapping) or isinstance(value, (list, tuple, set, frozenset)):
        return json.dumps(_safe_value(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    if isinstance(value, bytes):
        return "base64:" + base64.b64encode(value).decode("ascii")
    return str(value)


def _formula_like_text(value: str) -> bool:
    stripped = value.lstrip("".join(chr(code) for code in range(33)) + "\u00a0\u1680\u2000\u2001\u2002\u2003\u2004\u2005\u2006\u2007\u2008\u2009\u200a\u2028\u2029\u202f\u205f\u3000")
    return bool(stripped) and stripped[0] in "=+-@"


def _snapshot_audit(record: RunRecord) -> dict[str, Any]:
    metadata = record.metadata()
    try:
        details = record.snapshot_details()
    except Exception as exc:  # Report unavailability; never reconstruct it.
        details = {}
        details_error = f"{type(exc).__name__}: {exc}"
    else:
        details_error = None
    safe_details = {key: value for key, value in details.items() if key != "raw_sources"}
    return {
        "identity": {
            "run_id": record.run_id, "generation": record.generation,
            "revision_id": record.revision_id, "config_fingerprint": record.config_fingerprint,
            "phase": record.phase, "record_status": record.status, "created_at": record.created_at,
        },
        "record_metadata": _safe_value(metadata),
        "snapshot_details": _safe_value(safe_details) if safe_details else None,
        "snapshot_details_unavailable_reason": details_error,
        "logs": list(record.logs()),
        "raw_source_bytes_included": False,
    }


def _request_audit(record: RunRecord, request: TableExportRequest, frame: pd.DataFrame,
                   columns: tuple[str, ...], review_columns: tuple[str, ...]) -> dict[str, Any]:
    positions = _selection_positions(request, frame)
    return {
        "schema": "pyeyedata.table-export-audit/v1",
        **_snapshot_audit(record),
        "export": {
            "table_name": request.table_name, "format": request.file_format,
            "scope": request.scope, "filter_description": request.filter_description,
            "source_table_fingerprint": request.table_fingerprint,
            "source_row_count": len(frame), "export_row_count": len(positions),
            "selected_source_positions": list(request.source_positions),
            "selected_source_row_ids": list(request.source_row_ids),
            "columns": list(columns), "automatically_included_review_columns": list(review_columns),
            "schema": [{"position": index, "name": str(name), "dtype": str(dtype)}
                       for index, (name, dtype) in enumerate(zip(frame.columns, frame.dtypes))],
            "named_reference_tables": list(request.reference_tables),
            "csv_formula_escape": "Apostrophe prefix on formula-like string cells and headers after leading whitespace/control characters; original text and header mapping are recorded below. Numeric negatives remain numeric and unchanged.",
            "xlsx_text_policy": "User strings are stored with explicit string cell type; no formulas/macros are written.",
            "xlsx_limits": {"rows_including_header": XLSX_MAX_ROWS,
                            "columns": XLSX_MAX_COLUMNS, "cell_text_characters": XLSX_MAX_CELL_TEXT},
        },
    }


def _iter_output_rows(frame: pd.DataFrame, positions: Sequence[int] | range,
                      columns: tuple[str, ...]):
    by_name = {str(name): index for index, name in enumerate(frame.columns)}
    missing = [column for column in columns if column not in by_name]
    if missing:
        raise ExportIdentityError("output columns no longer match the source table: " + ", ".join(missing))
    column_positions = [by_name[column] for column in columns]
    for source_position in positions:
        source_row = frame.iloc[int(source_position)]
        row_id = source_row_identity(frame, int(source_position))
        yield int(source_position), row_id, tuple(source_row.iloc[column] for column in column_positions)


def preview_table_export(record: RunRecord, request: TableExportRequest) -> ExportPreview:
    """Pre-count exact rows/columns/sheets and safety transformations before writing."""
    frame = _verify_request(record, request)
    columns, review_columns = _output_columns(frame, request.columns)
    positions = _selection_positions(request, frame)
    row_count = len(positions)
    formula_count = 0
    over_limit = 0
    if request.file_format == "xlsx":
        for _source_position, _row_id, values in _iter_output_rows(frame, positions, columns):
            for value in values:
                text = _to_text(value)
                if len(text) > XLSX_MAX_CELL_TEXT:
                    over_limit += 1
                if isinstance(value, str) and _formula_like_text(value):
                    formula_count += 1
        for name in request.reference_tables:
            ref_frame = record.table_copy(name)
            ref_headers = [str(column) for column in ref_frame.columns]
            for row in ref_frame.itertuples(index=False, name=None):
                for header, value in zip(ref_headers, row):
                    cell_value = _excel_value(value, header)
                    if isinstance(cell_value, str) and _formula_like_text(cell_value):
                        formula_count += 1
                    if len(_to_text(cell_value)) > XLSX_MAX_CELL_TEXT:
                        over_limit += 1
        main_sheets = max(1, math.ceil(row_count / _XLSX_DATA_ROWS))
        formula_count += sum(_formula_like_text(column) for column in columns) * main_sheets
        ref_sheets = 0
        for name in request.reference_tables:
            ref_frame = record.table_copy(name)
            reference_sheets = max(1, math.ceil(len(ref_frame) / _XLSX_DATA_ROWS))
            ref_sheets += reference_sheets
            formula_count += sum(_formula_like_text(str(column)) for column in ref_frame.columns) * reference_sheets
        worksheets = main_sheets + ref_sheets + 3  # export_meta, formula_audit, long_text
        warnings: list[str] = []
        if row_count > _XLSX_DATA_ROWS and request.xlsx_overflow == "reject":
            warnings.append(f"需要 {main_sheets} 个主数据工作表；当前请求未允许拆分工作表。")
        if len(columns) > XLSX_MAX_COLUMNS:
            warnings.append(f"输出 {len(columns):,} 列超过 Excel 上限 {XLSX_MAX_COLUMNS:,}。")
        if over_limit:
            warnings.append(f"{over_limit:,} 个超长文本将以分段 LongText 工作表保存，并通过单元格引用还原。")
        overlong_headers = [column for column in columns if len(column) > XLSX_MAX_CELL_TEXT]
        for name in request.reference_tables:
            overlong_headers.extend(
                str(column) for column in record.table_copy(name).columns
                if len(str(column)) > XLSX_MAX_CELL_TEXT
            )
        if overlong_headers:
            warnings.append(
                f"{len(overlong_headers):,} 个列标题超过 Excel 的 {XLSX_MAX_CELL_TEXT:,} 字符限制；"
                "此 XLSX 请求会被拒绝以保留完整列名。"
            )
        return ExportPreview(
            record.run_id, record.revision_id, request.table_name, request.file_format, request.scope,
            row_count, len(frame), request.columns or tuple(str(column) for column in frame.columns),
            columns, review_columns, request.reference_tables, worksheets,
            row_count > _XLSX_DATA_ROWS, over_limit, formula_count, tuple(warnings),
        )
    formula_count = sum(_formula_like_text(column) for column in columns)
    for _source_position, _row_id, values in _iter_output_rows(frame, positions, columns):
        formula_count += sum(isinstance(value, str) and _formula_like_text(value) for value in values)
    return ExportPreview(
        record.run_id, record.revision_id, request.table_name, request.file_format, request.scope,
        row_count, len(frame), request.columns or tuple(str(column) for column in frame.columns),
        columns, review_columns, request.reference_tables, 0, False, 0, formula_count, (),
    )


def _atomic_temp_path(destination: Path, suffix: str = ".tmp") -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temp_name = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=suffix, dir=destination.parent)
    os.close(handle)
    return Path(temp_name)


def _write_json_temp(path: Path, payload: Mapping[str, Any]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(payload, stream, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False, default=str)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())


def _typed_reference_value(value: Any) -> Any:
    if _missing(value):
        return {"$missing": type(value).__name__}
    if isinstance(value, np.generic):
        return _typed_reference_value(value.item())
    if isinstance(value, (pd.Timestamp, datetime, date)):
        return {"$datetime": value.isoformat()}
    if isinstance(value, Decimal):
        return {"$decimal": str(value)}
    if isinstance(value, bytes):
        return {"$bytes_base64": base64.b64encode(value).decode("ascii")}
    if isinstance(value, float) and not math.isfinite(value):
        return {"$float": repr(value)}
    if isinstance(value, Mapping):
        return {str(key): _typed_reference_value(item) for key, item in sorted(value.items(), key=lambda pair: str(pair[0]))}
    if isinstance(value, (list, tuple)):
        return [_typed_reference_value(item) for item in value]
    if isinstance(value, set | frozenset):
        items = [_typed_reference_value(item) for item in value]
        return sorted(items, key=lambda item: json.dumps(item, ensure_ascii=False, sort_keys=True, default=str))
    return value


def _write_csv(
    record: RunRecord,
    request: TableExportRequest,
    frame: pd.DataFrame,
    columns: tuple[str, ...],
    csv_temp: Path,
    audit: dict[str, Any],
    reference_temp: Path | None,
    *,
    cancel: CancelCallback | None = None,
    progress: ProgressCallback | None = None,
) -> tuple[int, int]:
    positions = _selection_positions(request, frame)
    formula_cells: list[dict[str, Any]] = []
    rows_written = 0
    with csv_temp.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, dialect="excel", lineterminator="\n")
        header_mapping: list[dict[str, Any]] = []
        exported_headers: list[str] = []
        for column in columns:
            exported = "'" + column if _formula_like_text(column) else column
            exported_headers.append(exported)
            header_mapping.append({"canonical_name": column, "export_name": exported,
                                   "escaped": exported != column, "storage_type": "text"})
            if exported != column:
                formula_cells.append({
                    "source_position": "header", "source_row_id": "HEADER",
                    "column": column, "original_value": column,
                    "original_type": "str", "export_value": exported,
                })
        writer.writerow(exported_headers)
        audit["csv_header_mapping"] = header_mapping
        for source_position, row_id, values in _iter_output_rows(frame, positions, columns):
            _export_checkpoint(cancel, progress if rows_written % 128 == 0 else None,
                               0.1 + 0.6 * rows_written / max(1, len(positions)), "正在写入 CSV 行")
            output: list[Any] = []
            for column, value in zip(columns, values):
                cell_text = _to_text(value)
                if isinstance(value, str) and _formula_like_text(value):
                    formula_cells.append({
                        "source_position": source_position, "source_row_id": row_id,
                        "column": column, "original_value": value,
                        "original_type": type(value).__qualname__, "export_value": "'" + cell_text,
                    })
                    cell_text = "'" + cell_text
                output.append(cell_text)
            writer.writerow(output)
            rows_written += 1
        stream.flush()
        os.fsync(stream.fileno())
    audit["csv_formula_escaped_cells"] = formula_cells
    if reference_temp is not None:
        with reference_temp.open("w", encoding="utf-8", newline="\n") as stream:
            for reference_name in request.reference_tables:
                _export_checkpoint(cancel, progress, 0.75, "正在写入关联表")
                reference_frame = record.table_copy(reference_name)
                stream.write(json.dumps({
                    "kind": "named_table", "name": reference_name, "row_count": len(reference_frame),
                    "schema": [{"name": str(column), "dtype": str(dtype)}
                               for column, dtype in zip(reference_frame.columns, reference_frame.dtypes)],
                    "table_fingerprint": _table_fingerprint(reference_frame, cancel=cancel),
                }, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str) + "\n")
                for position, row in enumerate(reference_frame.itertuples(index=False, name=None)):
                    _export_checkpoint(cancel)
                    stream.write(json.dumps({
                        "kind": "row", "table": reference_name, "source_position": position,
                        "values": [_typed_reference_value(value) for value in row],
                    }, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str) + "\n")
            stream.flush()
            os.fsync(stream.fileno())
    return rows_written, len(formula_cells)


def _excel_value(value: Any, column_name: str) -> Any:
    if _missing(value):
        return None
    if isinstance(value, np.generic):
        value = value.item()
    if column_name in _ID_COLUMNS:
        return str(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.to_pydatetime() if isinstance(value, pd.Timestamp) else value
    if isinstance(value, date):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return "NaN" if math.isnan(value) else ("Infinity" if value > 0 else "-Infinity")
    if isinstance(value, bytes):
        return "base64:" + base64.b64encode(value).decode("ascii")
    if isinstance(value, Mapping) or isinstance(value, (list, tuple, set, frozenset)):
        return json.dumps(_safe_value(value), ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return value


def _write_excel_cell(ws: Any, value: Any) -> Any:
    from openpyxl.cell import WriteOnlyCell

    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _safe_sheet_name(raw: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", str(raw)).strip(" '\"") or "Table"
    base = base[:31]
    name = base
    index = 2
    while name.casefold() in {value.casefold() for value in used}:
        suffix = f"_{index}"
        name = base[:31 - len(suffix)] + suffix
        index += 1
    used.add(name)
    return name


def _append_metadata_sheet(ws: Any, metadata: Mapping[str, Any], *,
                           cancel: CancelCallback | None = None) -> None:
    ws.append([_write_excel_cell(ws, value) for value in ("field", "value_json")])
    for key, value in metadata.items():
        _export_checkpoint(cancel)
        text = json.dumps(_safe_value(value), ensure_ascii=False, sort_keys=True, default=str)
        if len(text) <= XLSX_MAX_CELL_TEXT:
            ws.append([_write_excel_cell(ws, str(key)), _write_excel_cell(ws, text)])
        else:
            parts = [text[index:index + _LONG_TEXT_CHUNK] for index in range(0, len(text), _LONG_TEXT_CHUNK)]
            for part_index, part in enumerate(parts, 1):
                _export_checkpoint(cancel)
                ws.append([_write_excel_cell(ws, f"{key} (part {part_index}/{len(parts)})"),
                           _write_excel_cell(ws, part)])


def _append_frame_to_xlsx(
    wb: Any,
    frame: pd.DataFrame,
    table_name: str,
    *,
    used_names: set[str],
    long_text_ws: Any,
    formula_audit_ws: Any,
    table_fingerprint: str,
    row_id_frame: pd.DataFrame | None = None,
    source_positions: Sequence[int] | range | None = None,
    cancel: CancelCallback | None = None,
    progress: ProgressCallback | None = None,
) -> tuple[list[str], int, int]:
    if frame.shape[1] > XLSX_MAX_COLUMNS:
        raise ExportLimitError(f"{table_name} has {frame.shape[1]:,} columns; Excel supports at most {XLSX_MAX_COLUMNS:,}.")
    num_sheets = max(1, math.ceil(len(frame) / _XLSX_DATA_ROWS))
    sheet_names: list[str] = []
    formula_count = 0
    long_count = 0
    for part_index in range(num_sheets):
        start = part_index * _XLSX_DATA_ROWS
        stop = min(len(frame), start + _XLSX_DATA_ROWS)
        _export_checkpoint(cancel, progress, 0.2 + 0.4 * start / max(1, len(frame)),
                           "正在创建 XLSX 工作表")
        raw_sheet = table_name if part_index == 0 else f"{table_name}_{part_index + 1}"
        ws = wb.create_sheet(_safe_sheet_name(raw_sheet, used_names))
        sheet_names.append(ws.title)
        headers = [str(column) for column in frame.columns]
        overlong_headers = [header for header in headers if len(header) > XLSX_MAX_CELL_TEXT]
        if overlong_headers:
            raise ExportLimitError(
                f"{table_name} has a column header longer than Excel's {XLSX_MAX_CELL_TEXT:,}-character limit; "
                "the schema header cannot be split without changing its meaning."
            )
        ws.append([_write_excel_cell(ws, header) for header in headers])
        for header in headers:
            if _formula_like_text(header):
                formula_audit_ws.append([
                    _write_excel_cell(formula_audit_ws, table_name),
                    _write_excel_cell(formula_audit_ws, "header"),
                    _write_excel_cell(formula_audit_ws, "HEADER"),
                    _write_excel_cell(formula_audit_ws, header),
                    _write_excel_cell(formula_audit_ws, "explicit-string-cell"),
                ])
                formula_count += 1
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{_column_letter(len(headers))}{max(1, stop - start + 1)}" if headers else None
        for col_num, header in enumerate(headers, 1):
            ws.column_dimensions[_column_letter(col_num)].width = min(max(len(header) + 2, 10), 48)
        for row_position in range(start, stop):
            _export_checkpoint(cancel, progress if row_position % 128 == 0 else None,
                               0.2 + 0.4 * row_position / max(1, len(frame)), "正在写入 XLSX 行")
            values = frame.iloc[row_position]
            source_position = (int(source_positions[row_position])
                               if source_positions is not None else row_position)
            identity_frame = row_id_frame if row_id_frame is not None else frame
            row_id = source_row_identity(identity_frame, source_position)
            cells: list[Any] = []
            for column, raw_value in zip(headers, values.tolist()):
                value = _excel_value(raw_value, column)
                if isinstance(value, str) and len(value) > XLSX_MAX_CELL_TEXT:
                    digest = hashlib.sha256(value.encode("utf-8")).hexdigest()[:20]
                    reference = f"LONGTEXT:{table_name}:{row_position}:{column}:{digest}"
                    chunks = [value[index:index + _LONG_TEXT_CHUNK]
                              for index in range(0, len(value), _LONG_TEXT_CHUNK)]
                    for chunk_index, chunk in enumerate(chunks, 1):
                        _export_checkpoint(cancel)
                        long_text_ws.append([
                            _write_excel_cell(long_text_ws, reference),
                            _write_excel_cell(long_text_ws, table_name),
                            _write_excel_cell(long_text_ws, source_position),
                            _write_excel_cell(long_text_ws, row_id),
                            _write_excel_cell(long_text_ws, column),
                            _write_excel_cell(long_text_ws, chunk_index),
                            _write_excel_cell(long_text_ws, len(chunks)),
                            _write_excel_cell(long_text_ws, chunk),
                        ])
                    value = f"[完整文本分段见 LongText: {reference}]"
                    long_count += 1
                if isinstance(value, str) and _formula_like_text(value):
                    formula_audit_ws.append([
                        _write_excel_cell(formula_audit_ws, table_name),
                        _write_excel_cell(formula_audit_ws, source_position),
                        _write_excel_cell(formula_audit_ws, row_id),
                        _write_excel_cell(formula_audit_ws, column),
                        _write_excel_cell(formula_audit_ws, "explicit-string-cell"),
                    ])
                    formula_count += 1
                cells.append(_write_excel_cell(ws, value))
            ws.append(cells)
    return sheet_names, formula_count, long_count


def _column_letter(number: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(max(1, number))


def _write_xlsx(
    record: RunRecord,
    request: TableExportRequest,
    frame: pd.DataFrame,
    columns: tuple[str, ...],
    temp_path: Path,
    audit: dict[str, Any],
    *,
    cancel: CancelCallback | None = None,
    progress: ProgressCallback | None = None,
) -> tuple[int, int, int, tuple[str, ...]]:
    from openpyxl import Workbook

    if len(columns) > XLSX_MAX_COLUMNS:
        raise ExportLimitError(f"输出 {len(columns):,} 列，超过 Excel 上限 {XLSX_MAX_COLUMNS:,}；未导出或截断。")
    row_count = len(_selection_positions(request, frame))
    if row_count > _XLSX_DATA_ROWS and request.xlsx_overflow == "reject":
        sheets_needed = math.ceil(row_count / _XLSX_DATA_ROWS)
        raise ExportLimitError(
            f"输出 {row_count:,} 行数据，需要 {sheets_needed} 个工作表（每表最多 {_XLSX_DATA_ROWS:,} 行数据）；"
            "请在预览中明确选择按工作表拆分或改用 CSV。"
        )
    overlong_main_headers = [column for column in columns if len(column) > XLSX_MAX_CELL_TEXT]
    if overlong_main_headers:
        raise ExportLimitError(
            f"{request.table_name} has a column header longer than Excel's {XLSX_MAX_CELL_TEXT:,}-character limit; "
            "the schema header cannot be split without changing its meaning."
        )
    for reference_name in request.reference_tables:
        _export_checkpoint(cancel)
        reference_frame = record.table_copy(reference_name)
        if reference_frame.shape[1] > XLSX_MAX_COLUMNS:
            raise ExportLimitError(
                f"{reference_name} has {reference_frame.shape[1]:,} columns; "
                f"Excel supports at most {XLSX_MAX_COLUMNS:,}."
            )
        overlong_headers = [str(column) for column in reference_frame.columns
                            if len(str(column)) > XLSX_MAX_CELL_TEXT]
        if overlong_headers:
            raise ExportLimitError(
                f"{reference_name} has a column header longer than Excel's {XLSX_MAX_CELL_TEXT:,}-character limit; "
                "the schema header cannot be split without changing its meaning."
            )
    wb = Workbook(write_only=True)
    try:
        used_names: set[str] = set()
        meta_ws = wb.create_sheet(_safe_sheet_name("export_meta", used_names))
        formula_ws = wb.create_sheet(_safe_sheet_name("formula_audit", used_names))
        formula_ws.append([_write_excel_cell(formula_ws, value) for value in
                           ("table", "source_position", "source_row_id", "column", "storage_policy")])
        long_ws = wb.create_sheet(_safe_sheet_name("LongText", used_names))
        long_ws.append([_write_excel_cell(long_ws, value) for value in
                        ("reference", "table", "source_position", "source_row_id", "column", "part", "parts", "text")])

        main_sheets: list[str] = []
        formula_count = 0
        long_count = 0
        positions = _selection_positions(request, frame)
        selected_positions = range(len(frame)) if request.scope == "full" else request.source_positions
        selected_frame = frame if request.scope == "full" else frame.iloc[list(selected_positions)].copy(deep=True)
        # Preserve request order for row subsets while retaining original row positions in metadata.
        by_name = {str(column): column for column in frame.columns}
        selected_columns = [by_name[column] for column in columns]
        selected_frame = selected_frame.loc[:, selected_columns]
        main_sheet_names, main_formula_count, main_long_count = _append_frame_to_xlsx(
            wb, selected_frame, request.table_name, used_names=used_names,
            long_text_ws=long_ws, formula_audit_ws=formula_ws,
            table_fingerprint=request.table_fingerprint, row_id_frame=frame,
            source_positions=selected_positions if request.scope != "full" else None,
            cancel=cancel, progress=progress,
        )
        main_sheets.extend(main_sheet_names)
        formula_count += main_formula_count
        long_count += main_long_count

        reference_sheet_names: list[str] = []
        for reference_name in request.reference_tables:
            _export_checkpoint(cancel, progress, 0.65, "正在写入 XLSX 关联表")
            reference_frame = record.table_copy(reference_name)
            ref_sheets, ref_formula, ref_long = _append_frame_to_xlsx(
                wb, reference_frame, reference_name, used_names=used_names,
                long_text_ws=long_ws, formula_audit_ws=formula_ws,
                table_fingerprint=_table_fingerprint(reference_frame, cancel=cancel),
                cancel=cancel,
            )
            reference_sheet_names.extend(ref_sheets)
            formula_count += ref_formula
            long_count += ref_long

        audit["xlsx_formula_like_text_cells"] = formula_count
        audit["xlsx_long_text_cells_spilled"] = long_count
        audit["xlsx_sheet_names"] = [*main_sheets, *reference_sheet_names, meta_ws.title,
                                    formula_ws.title, long_ws.title]
        audit["xlsx_long_text_policy"] = "Over-limit strings are split without truncation in the named LongText worksheet; source cells contain a stable reference."
        audit["xlsx_reference_policy"] = "Each requested named reference table is exported to its own worksheet(s) with the original schema."
        _append_metadata_sheet(meta_ws, audit, cancel=cancel)
        _export_checkpoint(cancel, progress, 0.8, "正在保存 XLSX；取消将等待当前保存结束")
        wb.save(temp_path)
        _export_checkpoint(cancel)
        rows_written = len(positions)
        return rows_written, formula_count, long_count, tuple(audit["xlsx_sheet_names"])
    finally:
        # Streaming worksheets own open XML generators and temporary XML files.
        # openpyxl only removes these on successful save; cancellation must also
        # close them. These paths are library-owned staging, never user outputs.
        for worksheet in wb.worksheets:
            try:
                if not worksheet.closed:
                    worksheet.close()
            finally:
                writer = worksheet._writer
                if writer is not None and os.path.exists(writer.out):
                    writer.cleanup()
        wb.close()


def _sha256(path: Path, *, cancel: CancelCallback | None = None) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            _export_checkpoint(cancel)
            digest.update(block)
    return digest.hexdigest()


def _commit_output_group(staged: Sequence[tuple[Path, Path]], *, overwrite: bool) -> None:
    """Promote a table and its sidecars together, restoring prior files on failure."""
    targets = [target for target, _temp in staged]
    existing = [target for target in targets if target.exists()]
    if existing and not overwrite:
        raise FileExistsError("export output already exists: " + ", ".join(str(path) for path in existing))
    backups: dict[Path, Path] = {}
    retain_backups = False
    try:
        if overwrite:
            import shutil

            for target in existing:
                backup = _atomic_temp_path(target, ".backup")
                try:
                    shutil.copy2(target, backup)
                except Exception:
                    backup.unlink(missing_ok=True)
                    raise
                backups[target] = backup
        promoted: list[Path] = []
        try:
            for target, temp in staged:
                os.replace(temp, target)
                promoted.append(target)
        except Exception as promotion_error:
            rollback_errors: list[str] = []
            for target in reversed(promoted):
                backup = backups.get(target)
                try:
                    if backup is not None:
                        os.replace(backup, target)
                        backups.pop(target, None)
                    else:
                        target.unlink(missing_ok=True)
                except OSError as rollback_error:
                    rollback_errors.append(f"{target}: {type(rollback_error).__name__}: {rollback_error}")
            if rollback_errors:
                retain_backups = True
                retained = ", ".join(f"{target} <- {backup}" for target, backup in backups.items())
                raise ExportError(
                    "output promotion failed and rollback was incomplete; recovery backups retained "
                    f"({retained}). Rollback errors: {'; '.join(rollback_errors)}"
                ) from promotion_error
            raise
        for backup in backups.values():
            try:
                backup.unlink(missing_ok=True)
            except OSError:
                # The output group has been committed; a hidden recovery copy is safer than reverting it.
                pass
    except Exception:
        if not retain_backups:
            for backup in backups.values():
                try:
                    backup.unlink(missing_ok=True)
                except OSError:
                    pass
        raise


def export_table(record: RunRecord, request: TableExportRequest, *,
                 cancel: CancelCallback | None = None,
                 progress: ProgressCallback | None = None) -> ExportResult:
    """Export stored values; cancel is a zero-argument predicate checked before commit.

    Progress receives (fraction, message). Promotion/rollback is non-interruptible;
    a returned receipt always represents committed output, even after late cancel.
    """
    _export_checkpoint(cancel, progress, 0.0, "正在核对导出内容")
    frame = _verify_request(record, request, cancel=cancel)
    columns, review_columns = _output_columns(frame, request.columns)
    destination = Path(request.destination)
    audit_path = destination.with_suffix(destination.suffix + ".audit.json")
    references_path = (destination.with_suffix(destination.suffix + ".references.jsonl")
                       if request.file_format == "csv" and request.reference_tables else None)
    outputs = [destination, audit_path] + ([references_path] if references_path else [])
    if not request.overwrite:
        existing = [str(path) for path in outputs if path.exists()]
        if existing:
            raise FileExistsError("export output already exists: " + ", ".join(existing))
    destination.parent.mkdir(parents=True, exist_ok=True)
    audit = _request_audit(record, request, frame, columns, review_columns)
    csv_temp = _atomic_temp_path(destination, ".csv.tmp")
    audit_temp = _atomic_temp_path(audit_path, ".json.tmp")
    reference_temp = _atomic_temp_path(references_path, ".jsonl.tmp") if references_path else None
    rows_written = 0
    formula_count = 0
    long_count = 0
    sheet_names: tuple[str, ...] = ()
    try:
        if request.file_format == "csv":
            rows_written, formula_count = _write_csv(
                record, request, frame, columns, csv_temp, audit, reference_temp,
                cancel=cancel, progress=progress,
            )
            audit["formula_like_text_cell_count"] = formula_count
            audit["references_sidecar"] = str(references_path) if references_path else None
        else:
            rows_written, formula_count, long_count, sheet_names = _write_xlsx(
                record, request, frame, columns, csv_temp, audit,
                cancel=cancel, progress=progress,
            )
            audit["worksheet_names"] = list(sheet_names)
            audit["formula_like_text_cell_count"] = formula_count
            audit["long_text_cells_spilled"] = long_count
        _export_checkpoint(cancel, progress, 0.9, "正在生成导出审计")
        audit["output_sha256"] = _sha256(csv_temp, cancel=cancel)
        _write_json_temp(audit_temp, audit)
        if not request.overwrite:
            # Recheck immediately before promotion to avoid silently clobbering a concurrent file.
            raced = [str(path) for path in outputs if path.exists()]
            if raced:
                raise FileExistsError("export output appeared during export: " + ", ".join(raced))
        staged: list[tuple[Path, Path]] = []
        if reference_temp is not None and references_path is not None:
            staged.append((references_path, reference_temp))
        staged.extend([(audit_path, audit_temp), (destination, csv_temp)])
        _export_checkpoint(cancel, progress, 0.98, "正在提交导出文件")
        _commit_output_group(staged, overwrite=request.overwrite)
    except Exception:
        for temp in (csv_temp, audit_temp, reference_temp):
            if temp is not None:
                try:
                    temp.unlink(missing_ok=True)
                except OSError:
                    pass
        raise
    return ExportResult(
        destination, audit_path, references_path, audit["output_sha256"], rows_written,
        columns, formula_count, long_count, sheet_names,
    )


__all__ = [
    "ExportCancelled", "ExportError", "ExportIdentityError", "ExportLimitError", "ExportPreview", "ExportResult",
    "TableExportRequest", "create_table_export_request", "export_table", "preview_table_export",
    "source_row_identity",
]
