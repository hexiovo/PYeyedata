"""Regenerate the versioned, ordinary-import templates from canonical contracts."""
from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import sys

from openpyxl import Workbook

PROJECT_ROOT = Path(__file__).resolve().parents[3]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from pyeyedata.contracts import (
    AOI_REQUIRED_COLUMNS,
    AOI_OPTIONAL_COLUMNS,
    CONTRACT_VERSION,
    SEGMENT_REQUIRED_COLUMNS,
    SEGMENT_OPTIONAL_COLUMNS,
    SIGNAL_OPTIONAL_COLUMNS,
    SIGNAL_REQUIRED_COLUMNS,
)
from pyeyedata.app_info import APP_VERSION


ROOT = Path(__file__).resolve().parent
RESOURCE_VERSION = "c03-resources/v3"
GUIDE_VERSION = "0.9.0"
TEMPLATE_VERSION = f"contract-{CONTRACT_VERSION}"

SIGNAL_COLUMNS = [
    "subject_id", "session_id", "time_s", "gaze_x", "gaze_y", "gaze_valid",
    "gaze_left_valid", "gaze_right_valid", "pupil_left", "pupil_right",
    "stimulus_id", "condition", "trial",
]
SEGMENT_COLUMNS = [
    "subject_id", "session_id", "segment_id", "segment_name", "start_s", "end_s",
    "condition", "trial", "stimulus_id",
]
AOI_COLUMNS = [
    "aoi_id", "stimulus_id", "shape", "aoi_name", "xmin", "ymin", "xmax", "ymax",
    "vertices", "unit", "coordinate_frame", "overlap_policy", "priority",
]
AOI_SCHEMA_REQUIRED = list(dict.fromkeys([*AOI_REQUIRED_COLUMNS, "aoi_name"]))
AOI_SCHEMA_OPTIONAL = [
    "xmin", "ymin", "xmax", "ymax", "vertices", "unit", "coordinate_frame",
    "overlap_policy", "priority",
]
AOI_DOMAIN_SCHEMA_VERSION = "2.0.0"
AOI_GEOMETRY_METHOD_VERSION = "static-aoi-geometry-2.0.0"
PERFORMANCE_COLUMNS = [
    "subject_id", "session_id", "condition", "trial", "accuracy", "reaction_time_s",
]

TEMPLATES = {
    "signals": {
        "columns": SIGNAL_COLUMNS,
        "rows": [[
            "SYNTHETIC_ONLY_TEMPLATE_P01", "SYNTHETIC_ONLY_TEMPLATE_S01", 0.0,
            500, 300, 1, 1, 1, 3.2, 3.1, "SYNTHETIC_ONLY_STIMULUS_01",
            "SYNTHETIC_ONLY_condition", 1,
        ]],
        "schema_required": list(SIGNAL_REQUIRED_COLUMNS),
        "schema_optional": list(SIGNAL_OPTIONAL_COLUMNS),
        "import_role": "signals",
        "normalizer": "normalize_signals",
    },
    "segments": {
        "columns": SEGMENT_COLUMNS,
        "rows": [[
            "SYNTHETIC_ONLY_TEMPLATE_P01", "SYNTHETIC_ONLY_TEMPLATE_S01",
            "SYNTHETIC_ONLY_TEMPLATE_SEG01", "合成演示区间", 0.0, 1.0,
            "SYNTHETIC_ONLY_condition", 1, "SYNTHETIC_ONLY_STIMULUS_01",
        ]],
        "schema_required": list(SEGMENT_REQUIRED_COLUMNS),
        "schema_optional": list(SEGMENT_OPTIONAL_COLUMNS),
        "import_role": "segments",
        "normalizer": "normalize_segments",
    },
    "aois": {
        "columns": AOI_COLUMNS,
        "rows": [[
            "SYNTHETIC_ONLY_TEMPLATE_AOI01", "SYNTHETIC_ONLY_STIMULUS_01",
            "polygon", "合成主要区域", "", "", "", "",
            "[[0,0],[100,0],[100,100],[0,100]]", "px", "synthetic_top_left", "exclusive", 1,
        ], [
            "SYNTHETIC_ONLY_TEMPLATE_AOI02", "SYNTHETIC_ONLY_STIMULUS_01",
            "polygon", "合成重叠区域", "", "", "", "",
            "[[50,50],[150,50],[150,150],[50,150]]", "px", "synthetic_top_left", "exclusive", 2,
        ]],
        "schema_required": AOI_SCHEMA_REQUIRED,
        "schema_optional": AOI_SCHEMA_OPTIONAL,
        "import_role": "aois",
        "normalizer": "normalize_aois",
        "domain_schema_version": AOI_DOMAIN_SCHEMA_VERSION,
        "geometry_method_version": AOI_GEOMETRY_METHOD_VERSION,
    },
    "performance": {
        "columns": PERFORMANCE_COLUMNS,
        "rows": [[
            "SYNTHETIC_ONLY_TEMPLATE_P01", "SYNTHETIC_ONLY_TEMPLATE_S01",
            "SYNTHETIC_ONLY_condition", 1, 1, 0.62,
        ]],
        "schema_required": [],
        "schema_optional": PERFORMANCE_COLUMNS,
        "import_role": "performance",
        "normalizer": "normalize_performance",
    },
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _refresh_release_resource_fields(index: dict[str, object], resource_root: Path) -> None:
    guide_path = "help/PYeyedata-user-guide-zh.pdf"
    licenses_root = resource_root / "licenses"
    license_paths = tuple(sorted(
        path.relative_to(resource_root).as_posix()
        for path in licenses_root.rglob("*")
        if path.is_file() and path.suffix.lower() in {".txt", ".md", ".json"}
    )) + ("fonts/OFL.txt",)
    missing = [
        relative for relative in (guide_path, *license_paths)
        if not (resource_root / relative).is_file()
    ]
    if missing:
        raise FileNotFoundError(f"Release resources are missing: {missing}")

    index["application_version"] = APP_VERSION
    help_entry = index.setdefault("help_guide", {})
    if not isinstance(help_entry, dict):
        raise TypeError("resource-index help_guide entry must be an object")
    help_entry.update({
        "status": "available",
        "application_version": APP_VERSION,
        "guide_version": GUIDE_VERSION,
        "resource_path": guide_path,
        "sha256": sha256(resource_root / guide_path),
    })
    source_manifest_path = licenses_root / "source-manifest.json"
    source_manifest = json.loads(source_manifest_path.read_text(encoding="utf-8"))
    index["licenses"] = {
        "status": source_manifest.get(
            "status", "offline license and source notices are available"
        ),
        "source_manifest": "licenses/source-manifest.json",
        "material_sha256": {
            relative: sha256(resource_root / relative) for relative in license_paths
        },
    }


def refresh_release_resource_index_only() -> None:
    """Refresh application-guide and license digests without rewriting templates."""
    resource_root = ROOT.parent
    index_path = resource_root / "resource-index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    _refresh_release_resource_fields(index, resource_root)
    index_path.write_text(
        json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def write_csv(path: Path, columns: list[str], rows: list[list[object]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\r\n")
        writer.writerow(columns)
        writer.writerows(rows)


def write_xlsx(path: Path, sheet_name: str, columns: list[str], rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.append(columns)
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def write_template_workbook(path: Path, role: str, specification: dict[str, object]) -> None:
    workbook = Workbook()
    empty = workbook.active
    empty.title = "empty_template"
    empty.append(specification["columns"])
    example = workbook.create_sheet("example")
    example.append(specification["columns"])
    for row in specification["rows"]:
        example.append(row)
    fields = workbook.create_sheet("fields")
    field_columns = ["canonical_field", "type", "unit", "status", "说明（中文）", "Description (English)"]
    fields.append(field_columns)
    for row in FIELD_ROWS[role]:
        fields.append(row)
    workbook.save(path)


FIELD_ROWS = {
    "signals": [
        ["subject_id", "text", "study ID", "required", "主体身份，必须显式提供。", "Explicit subject identity."],
        ["session_id", "text", "study ID", "required", "会话或录制身份，必须显式提供。", "Explicit session or recording identity."],
        ["time_s", "finite numeric", "s after import confirmation", "required", "确认真实时间单位、原点、时区/时钟和精度。", "Confirm source unit, origin, timezone/clock, and precision."],
        ["gaze_x", "finite numeric", "declared coordinate unit", "optional", "注视点横坐标。", "Gaze point horizontal coordinate."],
        ["gaze_y", "finite numeric", "declared coordinate unit", "optional", "注视点纵坐标。", "Gaze point vertical coordinate."],
        ["gaze_valid", "boolean code", "none", "optional", "源注视点有效性；自定义代码需显式映射。", "Source point validity; map custom labels explicitly."],
        ["gaze_left_valid", "boolean code", "none", "optional", "左眼注视有效性。", "Left-eye gaze validity."],
        ["gaze_right_valid", "boolean code", "none", "optional", "右眼注视有效性。", "Right-eye gaze validity."],
        ["pupil_left", "finite numeric", "source unit or unknown", "optional", "左眼原始瞳孔值；单位须说明。", "Raw left pupil value; declare its unit."],
        ["pupil_right", "finite numeric", "source unit or unknown", "optional", "右眼原始瞳孔值；单位须说明。", "Raw right pupil value; declare its unit."],
        ["stimulus_id", "text ID", "none", "optional", "与刺激定义表明确匹配。", "Must match an explicit stimulus definition."],
        ["condition", "text category", "none", "optional", "研究条件，保留真实设计层级。", "Study condition; preserve the design level."],
        ["trial", "integer or text ID", "none", "optional", "试次身份，保留真实重复结构。", "Trial identity; preserve repeated structure."],
    ],
    "segments": [
        ["subject_id", "text ID", "none", "required in participant/session mode", "主体身份；共享互动须用显式成员映射。", "Subject identity; shared interaction requires explicit member mapping."],
        ["session_id", "text ID", "none", "required in participant/session mode", "会话身份；共享互动须用显式成员映射。", "Session identity; shared interaction requires explicit member mapping."],
        ["segment_id", "text ID", "none", "optional", "显式区间ID；未提供时由导入/合同生成稳定标识。", "Optional explicit interval ID; otherwise a stable ID is generated."],
        ["segment_name", "text", "none", "required", "人工定义区间名称。", "Human-defined segment name."],
        ["start_s", "finite numeric", "s after import confirmation", "required", "左闭右开区间起点；确认真实时间单位。", "Left-closed interval start; confirm source time unit."],
        ["end_s", "finite numeric", "s after import confirmation", "required", "右边界；必须大于start_s。", "Right boundary; must exceed start_s."],
        ["condition", "text category", "none", "optional", "真实研究条件。", "Actual study condition."],
        ["trial", "integer or text ID", "none", "optional", "真实试次身份。", "Actual trial identity."],
        ["stimulus_id", "text ID", "none", "optional", "与信号中的刺激身份匹配。", "Matches stimulus identity in signals."],
    ],
    "aois": [
        ["aoi_id", "text ID", "none", "required input", "稳定AOI身份；用于连接和结果键。", "Stable AOI identity used for joins and result keys."],
        ["stimulus_id", "text ID", "none", "required input", "必须与刺激定义明确匹配。", "Must match an explicit stimulus definition."],
        ["shape", "enum", "none", "required input", "新表使用rectangle或polygon。", "New tables use rectangle or polygon."],
        ["aoi_name", "nonblank text", "none", "required input", "人类可读来源名称；保留中文，不替代aoi_id。", "Human-readable source name; preserved as supplied and never replaces aoi_id."],
        ["xmin", "finite numeric", "declared coordinate unit", "rectangle input", "矩形左边界；须满足xmax > xmin。", "Rectangle left bound; require xmax > xmin."],
        ["ymin", "finite numeric", "declared coordinate unit", "rectangle input", "矩形上边界；须满足ymax > ymin。", "Rectangle top bound; require ymax > ymin."],
        ["xmax", "finite numeric", "declared coordinate unit", "rectangle input", "矩形右边界。", "Rectangle right bound."],
        ["ymax", "finite numeric", "declared coordinate unit", "rectangle input", "矩形下边界。", "Rectangle bottom bound."],
        ["vertices", "JSON text Nx2", "declared coordinate unit", "polygon input", "写成[[x,y],...]；3至1024个不同顶点；不自动修复。", "Use [[x,y],...]; 3–1,024 distinct vertices; no automatic repair."],
        ["unit", "text", "coordinate unit", "optional input", "显式声明坐标单位；不会自动缩放。", "Declare coordinate units; no automatic scaling."],
        ["coordinate_frame", "text", "coordinate frame", "optional input", "显式声明坐标frame；不会推断坐标变换。", "Declare the coordinate frame; no inferred transform."],
        ["overlap_policy", "enum", "none", "optional input", "exclusive须配唯一priority；multi_label保留所有命中。", "exclusive needs unique priorities; multi_label retains all hits."],
        ["priority", "finite numeric", "none", "conditional input", "exclusive模式的命中AOI须有唯一priority；数值较小者优先。", "Exclusive hits need unique priorities; the lowest value wins."],
        ["aoi_name_status", "audit text", "none", "generated output", "来源名称状态；旧记录缺名只能走显式审计兼容路径。", "Name provenance; legacy missing names require explicit audited compatibility."],
        ["aoi_domain_schema_version", "version text", "none", "generated output", "AOI域schema版本2.0.0；不要写入源表。", "AOI domain schema 2.0.0; do not enter in source tables."],
        ["aoi_geometry_version", "version text", "none", "generated output", "静态几何方法版本；不要写入源表。", "Static geometry method version; do not enter in source tables."],
        ["aoi_boundary_policy", "audit text", "none", "generated output", "边与顶点按closed边界命中。", "Edges and vertices are included under closed boundaries."],
        ["aoi_input_closure", "audit text", "none", "generated output", "记录顶点输入使用隐式或显式闭合。", "Records whether the input polygon closure was implicit or explicit."],
        ["aoi_geometry_fingerprint", "hash text", "none", "generated output", "规范几何指纹；不含名称或来源路径。", "Canonical geometry fingerprint; excludes name and source path."],
        ["aoi_source", "structured audit", "none", "generated output", "保存来源文件、sheet、行和字段映射。", "Records source file, sheet, row, and field mapping."],
        ["aoi_geometry_raw", "raw audit", "none", "generated output", "保留原始顶点文本或矩形边界值。", "Retains original vertex text or rectangle bounds."],
        ["aoi_geometry_audit", "structured audit", "none", "generated output", "记录几何检查、容量与闭合结果。", "Records geometry checks, capacity limits, and closure results."],
    ],
    "performance": [
        ["subject_id", "text ID", "none", "candidate join key", "主体键须由设计显式选择。", "Select the subject key explicitly from the design."],
        ["session_id", "text ID", "none", "candidate join key", "会话键须由设计显式选择。", "Select the session key explicitly from the design."],
        ["condition", "text category", "none", "candidate join key", "重复条件数据需要时作为显式连接键。", "Use as an explicit join key when required by repeated conditions."],
        ["trial", "integer or text ID", "none", "candidate join key", "试次键须满足明确的一对一连接。", "Must satisfy the declared one-to-one join."],
        ["accuracy", "binary", "0/1", "example column", "二元表现，不可强作连续结果。", "Binary target; do not force it into a continuous method."],
        ["reaction_time_s", "finite numeric", "seconds", "example column", "连续反应时示例；确认原始单位。", "Continuous RT example; confirm the source unit."],
    ],
}


def refresh_aoi_resources_only() -> None:
    """Refresh AOI CSVs and AOI manifest entries without touching other resources.

    The AOI workbook is authored separately with @oai/artifact-tool. This mode
    requires that workbook to exist so it can record its exact hash and size.
    """
    role = "aois"
    specification = TEMPLATES[role]
    stem = f"{role}-{TEMPLATE_VERSION}"
    empty_csv = ROOT / f"{stem}.csv"
    example_csv = ROOT / f"{role}-example-{TEMPLATE_VERSION}.csv"
    workbook_path = ROOT / f"{stem}.xlsx"
    if not workbook_path.is_file():
        raise FileNotFoundError(
            "Create aois-contract-1.3.0.xlsx with @oai/artifact-tool before refreshing AOI manifests."
        )

    write_csv(empty_csv, specification["columns"], [])
    write_csv(example_csv, specification["columns"], specification["rows"])
    validation_status = "Resource-owner checks passed for public AOI normalization and static geometry audit on bundled examples; independent C05 acceptance is not claimed."
    records = [
        {
            "path": empty_csv.name, "role": role, "kind": "empty_template", "format": "csv",
            "sha256": sha256(empty_csv), "byte_size": empty_csv.stat().st_size,
            "data_rows": 0, "import_role": specification["import_role"],
            "normalizer": specification["normalizer"], "source_schema_required": specification["schema_required"],
            "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
            "label": "Header-only empty template", "status": "available",
        },
        {
            "path": example_csv.name, "role": role, "kind": "example", "format": "csv",
            "sha256": sha256(example_csv), "byte_size": example_csv.stat().st_size,
            "data_rows": len(specification["rows"]), "import_role": specification["import_role"],
            "normalizer": specification["normalizer"], "source_schema_required": specification["schema_required"],
            "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
            "label": "SYNTHETIC_ONLY; replace or remove before real use", "status": "available",
        },
        {
            "path": workbook_path.name, "role": role, "kind": "three_sheet_workbook", "format": "xlsx",
            "sha256": sha256(workbook_path), "byte_size": workbook_path.stat().st_size,
            "worksheets": ["empty_template", "example", "fields"],
            "data_rows_by_sheet": {"empty_template": 0, "example": len(specification["rows"]), "fields": len(FIELD_ROWS[role])},
            "import_role": specification["import_role"], "normalizer": specification["normalizer"],
            "source_schema_required": specification["schema_required"],
            "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
            "label": "Choose empty_template or example in the GUI; fields is documentation and never an import sheet",
            "status": "available",
        },
    ]
    for record in records:
        record["domain_schema_version"] = specification["domain_schema_version"]
        record["geometry_method_version"] = specification["geometry_method_version"]
        record["validation_status"] = validation_status

    manifest_path = ROOT / "template-manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["resource_version"] = RESOURCE_VERSION
    manifest["generator_sha256"] = sha256(Path(__file__).resolve())
    manifest["template_files"] = [
        item for item in manifest["template_files"] if item.get("role") != role
    ] + records
    manifest["aoi_contract"] = {
        "domain_schema_version": AOI_DOMAIN_SCHEMA_VERSION,
        "geometry_method_version": AOI_GEOMETRY_METHOD_VERSION,
        "boundary_policy": "closed",
        "validation_status": validation_status,
    }
    manifest["notes"] = [
        "For signals, segments, and performance, each XLSX workbook contains empty_template (header only), example (SYNTHETIC_ONLY row), and fields (documentation); GUI users must select empty_template or example explicitly, never fields.",
        "CSV empty templates and synthetic examples are separate files. Empty templates can be read but cannot be normalized until populated.",
        "AOI domain schema 2.0.0 adds required source aoi_name and JSON-text polygon vertices; its method metadata is static-aoi-geometry-2.0.0.",
        "Time and coordinate semantics are explicit import-dialog choices; the files do not imply device verification.",
        "AOI boundary semantics are closed. Capacity limits, closure rules, overlap policies, and legacy-name migration are documented in the bilingual field guide.",
        "Resource-owner AOI example checks passed public normalization and geometry audit; this is not independent C05 acceptance or scientific/device/physiology evidence.",
    ]
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    resource_root = ROOT.parent
    index_path = resource_root / "resource-index.json"
    index = json.loads(index_path.read_text(encoding="utf-8"))
    index["resource_version"] = RESOURCE_VERSION
    aoi_entry = index["roles"][role]
    aoi_entry.update({
        "workbook": f"templates/{stem}.xlsx",
        "empty_template": {
            "csv": f"templates/{stem}.csv",
            "xlsx": f"templates/{stem}.xlsx",
            "xlsx_sheet": "empty_template",
        },
        "example": {
            "csv": f"templates/{role}-example-{TEMPLATE_VERSION}.csv",
            "xlsx": f"templates/{stem}.xlsx",
            "xlsx_sheet": "example",
            "synthetic_label": "SYNTHETIC_ONLY",
        },
        "field_sheet": {"workbook": f"templates/{stem}.xlsx", "sheet": "fields"},
        "workbook_sheet_instruction": "Explicitly select empty_template or example; the fields sheet is documentation and must not be imported.",
        "status": "AOI source resources are available; resource-owner normalization and geometry checks passed for the bundled example; independent C05 acceptance is not claimed.",
        "note": "The example uses JSON polygon vertices, explicit exclusive priorities, and declared coordinate metadata; its resource-owner normalization and geometry audit passed.",
        "domain_schema_version": AOI_DOMAIN_SCHEMA_VERSION,
        "geometry_method_version": AOI_GEOMETRY_METHOD_VERSION,
        "boundary_policy": "closed",
        "source_columns_required": specification["schema_required"],
        "source_columns_optional": specification["schema_optional"],
    })
    index["help_guide"]["resource_path"] = "help/PYeyedata-user-guide-zh.pdf"
    _refresh_release_resource_fields(index, resource_root)
    index_path.write_text(json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    if sys.argv[1:] == ["--release-index-only"]:
        refresh_release_resource_index_only()
        return
    if sys.argv[1:] == ["--aois-only"]:
        refresh_aoi_resources_only()
        return

    records = []
    for role, specification in TEMPLATES.items():
        stem = f"{role}-{TEMPLATE_VERSION}"
        empty_csv = ROOT / f"{stem}.csv"
        example_csv = ROOT / f"{role}-example-{TEMPLATE_VERSION}.csv"
        workbook_path = ROOT / f"{stem}.xlsx"
        write_csv(empty_csv, specification["columns"], [])
        write_csv(example_csv, specification["columns"], specification["rows"])
        if role == "aois":
            if not workbook_path.is_file():
                raise FileNotFoundError(
                    "Create aois-contract-1.3.0.xlsx with @oai/artifact-tool before generating template manifests."
                )
        else:
            write_template_workbook(workbook_path, role, specification)
        records.extend([
            {
                "path": empty_csv.name, "role": role, "kind": "empty_template", "format": "csv",
                "sha256": sha256(empty_csv), "byte_size": empty_csv.stat().st_size,
                "data_rows": 0, "import_role": specification["import_role"],
                "normalizer": specification["normalizer"], "source_schema_required": specification["schema_required"],
                "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
                "label": "Header-only empty template", "status": "available",
            },
            {
                "path": example_csv.name, "role": role, "kind": "example", "format": "csv",
                "sha256": sha256(example_csv), "byte_size": example_csv.stat().st_size,
                "data_rows": len(specification["rows"]), "import_role": specification["import_role"],
                "normalizer": specification["normalizer"], "source_schema_required": specification["schema_required"],
                "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
                "label": "SYNTHETIC_ONLY; replace or remove before real use", "status": "available",
            },
            {
                "path": workbook_path.name, "role": role, "kind": "three_sheet_workbook", "format": "xlsx",
                "sha256": sha256(workbook_path), "byte_size": workbook_path.stat().st_size,
                "worksheets": ["empty_template", "example", "fields"],
                "data_rows_by_sheet": {"empty_template": 0, "example": len(specification["rows"]), "fields": len(FIELD_ROWS[role])},
                "import_role": specification["import_role"], "normalizer": specification["normalizer"],
                "source_schema_required": specification["schema_required"],
                "source_schema_optional": specification["schema_optional"], "columns": specification["columns"],
                "label": "Choose empty_template or example in the GUI; fields is documentation and never an import sheet",
                "status": "available",
            },
        ])
        if role == "aois":
            for record in records[-3:]:
                record["domain_schema_version"] = specification["domain_schema_version"]
                record["geometry_method_version"] = specification["geometry_method_version"]
                record["validation_status"] = "Resource-owner checks passed for public AOI normalization and static geometry audit on bundled examples; independent C05 acceptance is not claimed."

    manifest = {
        "schema": "pyeyedata.template-manifest/v2",
        "resource_version": RESOURCE_VERSION,
        "template_version": TEMPLATE_VERSION,
        "contract_version": CONTRACT_VERSION,
        "template_files": records,
        "generator": "generate_templates.py",
        "generator_sha256": sha256(Path(__file__).resolve()),
        "field_guide": "field-guide-zh-en.md",
        "example_pages": "example-pages-zh-en.md",
        "import_route": "read_table(role=...), followed by the listed public normalizer; no template-only importer",
        "aoi_contract": {
            "domain_schema_version": AOI_DOMAIN_SCHEMA_VERSION,
            "geometry_method_version": AOI_GEOMETRY_METHOD_VERSION,
            "boundary_policy": "closed",
            "validation_status": "Resource-owner checks passed for public AOI normalization and static geometry audit on bundled examples; independent C05 acceptance is not claimed.",
        },
        "notes": [
            "For signals, segments, and performance, each XLSX workbook contains empty_template (header only), example (SYNTHETIC_ONLY row), and fields (documentation); GUI users must select empty_template or example explicitly, never fields.",
            "CSV empty templates and synthetic examples are separate files. Empty templates can be read but cannot be normalized until populated.",
            "AOI domain schema 2.0.0 adds required source aoi_name and JSON-text polygon vertices; its method metadata is static-aoi-geometry-2.0.0.",
            "Time and coordinate semantics are explicit import-dialog choices; the files do not imply device verification.",
            "AOI boundary semantics are closed. Capacity limits, closure rules, overlap policies, and legacy-name migration are documented in the bilingual field guide.",
            "Resource-owner AOI example checks passed public normalization and geometry audit; this is not independent C05 acceptance or scientific/device/physiology evidence.",
        ],
    }
    (ROOT / "template-manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    index = {
        "schema": "pyeyedata.resource-index/v1",
        "resource_version": RESOURCE_VERSION,
        "contract_version": CONTRACT_VERSION,
        "resource_path_api": "pyeyedata.resource_paths.resource_path(relative_path: str) -> ContextManager[Path]",
        "resource_root": "pyeyedata/resources/",
        "roles": {
            role: {
                "workbook": f"templates/{role}-{TEMPLATE_VERSION}.xlsx",
                "empty_template": {
                    "csv": f"templates/{role}-{TEMPLATE_VERSION}.csv",
                    "xlsx": f"templates/{role}-{TEMPLATE_VERSION}.xlsx",
                    "xlsx_sheet": "empty_template",
                },
                "example": {
                    "csv": f"templates/{role}-example-{TEMPLATE_VERSION}.csv",
                    "xlsx": f"templates/{role}-{TEMPLATE_VERSION}.xlsx",
                    "xlsx_sheet": "example",
                    "synthetic_label": "SYNTHETIC_ONLY",
                },
                "field_sheet": {
                    "workbook": f"templates/{role}-{TEMPLATE_VERSION}.xlsx",
                    "sheet": "fields",
                },
                "field_guide": "templates/field-guide-zh-en.md",
                "example_page": "templates/example-pages-zh-en.md",
                "workbook_sheet_instruction": "Explicitly select empty_template or example; the fields sheet is documentation and must not be imported.",
                "status": "available" if role != "aois" else "AOI resources available; resource-owner example normalization and geometry audit passed; independent C05 acceptance is not claimed.",
                "note": "Empty template is header-only; example file contains one SYNTHETIC_ONLY row." if role != "aois" else "AOI inputs require aoi_name. The example uses JSON polygon vertices, explicit exclusive priorities, and declared coordinate metadata; resource-owner normalization and geometry audit passed.",
                **({
                    "domain_schema_version": AOI_DOMAIN_SCHEMA_VERSION,
                    "geometry_method_version": AOI_GEOMETRY_METHOD_VERSION,
                    "boundary_policy": "closed",
                    "source_columns_required": AOI_SCHEMA_REQUIRED,
                    "source_columns_optional": AOI_SCHEMA_OPTIONAL,
                } if role == "aois" else {}),
            }
            for role in TEMPLATES
        },
        "template_manifest": "templates/template-manifest.json",
        "synthetic_demos": {
            "manifest": "demo/demo-manifest.json",
            "generator": "demo/generate_demo.py",
            "scenario_ids": ["A", "B", "C", "D"],
            "label": "SYNTHETIC_ONLY",
            "note": "The manifest lists exact files and hashes. These are software-path fixtures only; not participant/device/physiology evidence.",
        },
        "help_guide": {
            "status": "available",
            "resource_path": "help/PYeyedata-user-guide-zh.pdf",
        },
        "fonts": {
            "path": "fonts/PYeyedataCJK-Regular.ttf",
            "license": "fonts/OFL.txt",
            "source_manifest": "fonts/font-source.json",
            "usage": "Keep the resource_path context open through registration and the complete render/export.",
        },
    }
    resource_root = ROOT.parent
    _refresh_release_resource_fields(index, resource_root)
    (resource_root / "resource-index.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


if __name__ == "__main__":
    main()
